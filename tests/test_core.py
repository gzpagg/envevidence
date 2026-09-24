import csv
import io
import json
from copy import deepcopy

import httpx
import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter

from envevidence.demo import DemoProvider, assets, demo_project
from envevidence.exporting import csv_bytes, xlsx_bytes
from envevidence.models import SourceRef
from envevidence.parsing import ParseError, parse_pdf
from envevidence.pipeline import run_project
from envevidence.providers import APIProvider, ProviderError, input_text
from envevidence.storage import ProjectStore
from envevidence.verification import locate, revise_field, validate_extraction


@pytest.fixture
def project():
    return demo_project()


@pytest.fixture
def populated(project, tmp_path):
    return run_project(project, ProjectStore(tmp_path), DemoProvider())


def find_field(project, experiment, key):
    return next(
        f for e in project.experiments if e.label == experiment for f in e.fields if f.key == key
    )


def test_end_to_end_expected_separate_conditions_and_metrics(populated):
    expected = json.loads(assets().joinpath("expected.json").read_text(encoding="utf-8"))
    assert len(populated.experiments) == expected["experiments"]
    for name, values in expected["values"].items():
        for key, value in values.items():
            assert find_field(populated, name, key).value == value
    assert find_field(populated, "Trial A", "removal_efficiency").unit == "%"
    assert (
        find_field(populated, "Trial C", "mineralization_efficiency").original.status == "not_found"
    )
    assert all(f.review_status == "pending" for e in populated.experiments for f in e.fields)
    assert all(s.matched for e in populated.experiments for f in e.fields for s in f.sources)


def test_supplement_provenance(populated):
    for trial in ["Trial A", "Trial B"]:
        field = find_field(populated, trial, "dose")
        assert field.sources[0].filename == "synthetic_supplement.pdf"
        assert field.sources[0].page == 1
        assert field.sources[0].matched


def test_fabricated_source_or_quote_is_not_verified(project):
    study = project.studies[0]
    block = study.documents[0].blocks[0]
    result = locate(
        [
            SourceRef(block_id=block.id, quote="A fabricated result of 999%."),
            SourceRef(block_id="../foreign-document:p1", quote=block.text),
        ],
        study,
    )
    assert not any(s.matched for s in result)
    assert result[1].filename is None


def test_whitespace_normalization_keeps_numbers_distinct(project):
    study = project.studies[0]
    block = study.documents[0].blocks[1]
    quote = "Trial A: acetaminophen removal was 85%; TOC removal was 20%."
    assert locate([SourceRef(block_id=block.id, quote=quote.replace(" ", "\n"))], study)[0].matched
    assert not locate([SourceRef(block_id=block.id, quote=quote.replace("85", "95"))], study)[
        0
    ].matched


def test_missing_supplement_cannot_validate_dose(project):
    study = project.studies[0]
    raw = DemoProvider().extract(study, project.fields).extraction
    study.documents = [d for d in study.documents if d.role == "main"]
    output = validate_extraction(raw, study, project.fields)
    dose = next(f for f in output[0].fields if f.key == "dose")
    assert not dose.sources[0].matched
    assert dose.issues


def test_missing_unit_is_flagged(project):
    raw = DemoProvider().extract(project.studies[0], project.fields).extraction
    next(f for f in raw.experiments[0].fields if f.key == "removal_efficiency").unit = None
    output = validate_extraction(raw, project.studies[0], project.fields)
    assert any("缺少单位" in issue for f in output[0].fields for issue in f.issues)


@pytest.mark.parametrize("change", ["missing", "duplicate", "unknown", "not_found_with_value"])
def test_inconsistent_response_rejected(project, change):
    raw = DemoProvider().extract(project.studies[0], project.fields).extraction
    if change == "missing":
        raw.experiments[0].fields.pop()
    elif change == "duplicate":
        raw.experiments[0].fields.append(raw.experiments[0].fields[0])
    elif change == "unknown":
        raw.experiments[0].fields[0].key = "not_requested"
    else:
        raw.experiments[0].fields[0].status = "not_found"
    with pytest.raises(ValueError):
        validate_extraction(raw, project.studies[0], project.fields)


def test_revision_keeps_original_and_roundtrips(populated, tmp_path):
    field = find_field(populated, "Trial A", "removal_efficiency")
    original = deepcopy(field.original)
    sources = [SourceRef(block_id=s.block_id, quote=s.quote) for s in field.sources]
    revise_field(
        field,
        populated.studies[0],
        value="85",
        unit="%",
        status="verified",
        reviewer="Tester",
        note="Checked A, not B, against the original page.",
        sources=sources,
    )
    store = ProjectStore(tmp_path)
    store.save(populated)
    loaded = store.load(populated.id)
    revised = find_field(loaded, "Trial A", "removal_efficiency")
    assert revised.original == original
    assert revised.review_status == "verified"
    assert revised.revisions[0].reviewer == "Tester"


def test_cannot_confirm_without_source(populated):
    field = find_field(populated, "Trial C", "mineralization_efficiency")
    with pytest.raises(ValueError):
        revise_field(
            field,
            populated.studies[0],
            value="90",
            unit="%",
            status="verified",
            note="guess",
            reviewer="Tester",
            sources=[],
        )
    assert not field.revisions


def test_interruption_preserves_first_study_and_resume_skips_it(project, tmp_path):
    second = project.studies[0].model_copy(deep=True)
    second.id, second.name = "second-study", "Second study"
    project.studies.append(second)
    store = ProjectStore(tmp_path)

    class Interrupted:
        def extract(self, study, fields):
            if study.id == second.id:
                raise ProviderError("Simulated API interruption")
            return DemoProvider().extract(study, fields)

    run_project(project, store, Interrupted())
    saved = store.load(project.id)
    assert len(saved.experiments) == 3
    assert [r.status for r in saved.runs] == ["completed", "failed"]
    run_project(saved, store, DemoProvider())
    assert len(saved.experiments) == 6
    assert [r.status for r in saved.runs] == ["completed", "failed", "completed"]


def test_unexpected_error_does_not_persist_secret(project, tmp_path):
    class BadProvider:
        def extract(self, *_):
            raise RuntimeError("secret-credential-and-document-text")

    store = ProjectStore(tmp_path)
    run_project(project, store, BadProvider())
    assert "secret-credential" not in store.path(project.id).read_text(encoding="utf-8")


def test_exports_and_formula_injection(populated):
    field = find_field(populated, "Trial A", "pollutant")
    field.original.value = '=HYPERLINK("https://invalid.example","click")'
    csv_rows = list(csv.DictReader(io.StringIO(csv_bytes(populated).decode("utf-8-sig"))))
    assert len(csv_rows) == 27
    assert csv_rows[0]["value"].startswith("'=")
    workbook = load_workbook(io.BytesIO(xlsx_bytes(populated)))
    assert workbook.sheetnames == ["Evidence", "Experiments", "Revisions", "Documents", "Runs"]
    assert workbook["Evidence"].max_row == 28
    assert workbook["Experiments"].max_row == 4
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)


def test_invalid_and_scanned_and_encrypted_pdfs_rejected():
    with pytest.raises(ParseError):
        parse_pdf(b"not a pdf", "bad.pdf")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(ParseError, match="扫描"):
        parse_pdf(stream.getvalue(), "scan.pdf")
    writer.encrypt("password")
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(ParseError, match="加密"):
        parse_pdf(stream.getvalue(), "encrypted.pdf")


def test_excel_handles_control_characters_and_marks_oversize_cells(populated):
    field = find_field(populated, "Trial A", "pollutant")
    field.original.value = "\x00" + "x" * 33000
    workbook = load_workbook(io.BytesIO(xlsx_bytes(populated)))
    cells = list(workbook["Evidence"].values)
    index = cells[0].index("value")
    value = cells[1][index]
    assert value.startswith("[U+0000]")
    assert "TRUNCATED" in value
    assert len(field.original.value) == 33001


def test_partial_page_warning_and_pdf_page_number_preserved():
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(assets().joinpath("synthetic_main.pdf").read_bytes()))
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_page(reader.pages[1])
    stream = io.BytesIO()
    writer.write(stream)
    result = parse_pdf(stream.getvalue(), "mixed.pdf")
    assert result.page_count == 2
    assert result.blocks[0].page == 2
    assert "第 1 页" in result.warnings[0]


def test_storage_rejects_traversal(tmp_path):
    with pytest.raises(ValueError):
        ProjectStore(tmp_path).load("../../secret")


def test_large_study_not_silently_truncated(project):
    project.studies[0].documents[0].blocks[0].text = "x" * 170_000
    with pytest.raises(ProviderError, match="未截断"):
        input_text(project.studies[0], project.fields)


@pytest.mark.parametrize("provider", ["openai", "anthropic"])
def test_api_contracts_with_mock_transport(project, provider):
    raw = DemoProvider().extract(project.studies[0], project.fields).extraction.model_dump_json()

    def handler(request):
        body = json.loads(request.content)
        assert body["model"] == "test-model"
        if provider == "openai":
            assert str(request.url) == "https://api.openai.com/v1/responses"
            assert body["store"] is False
            assert body["text"]["format"]["strict"] is True
            assert request.headers["Authorization"] == "Bearer test-key"
            payload = {
                "status": "completed",
                "output": [{"type": "message", "content": [{"type": "output_text", "text": raw}]}],
            }
        else:
            assert str(request.url) == "https://api.anthropic.com/v1/messages"
            assert body["output_config"]["format"]["type"] == "json_schema"
            assert request.headers["x-api-key"] == "test-key"
            payload = {"stop_reason": "end_turn", "content": [{"type": "text", "text": raw}]}
        payload["usage"] = {"input_tokens": 123, "output_tokens": 456}
        return httpx.Response(200, json=payload)

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = APIProvider(provider, "test-model", "test-key", client=client).extract(
            project.studies[0], project.fields
        )
    assert len(result.extraction.experiments) == 3
    assert result.usage["input_tokens"] == 123


@pytest.mark.parametrize("code", [400, 401, 403, 429, 500])
def test_provider_error_redacts_body(project, code):
    with httpx.Client(
        transport=httpx.MockTransport(lambda r: httpx.Response(code, text="SECRET"))
    ) as client:
        with pytest.raises(ProviderError) as exc:
            APIProvider("openai", "test", "key", client=client).extract(
                project.studies[0], project.fields
            )
        assert "SECRET" not in str(exc.value)


@pytest.mark.parametrize(
    "provider,payload",
    [
        ("openai", {"status": "incomplete", "output": []}),
        ("anthropic", {"stop_reason": "max_tokens", "content": []}),
        (
            "openai",
            {
                "status": "completed",
                "output": [{"type": "message", "content": [{"type": "refusal"}]}],
            },
        ),
        (
            "anthropic",
            {"stop_reason": "end_turn", "content": [{"type": "text", "text": "not-json"}]},
        ),
    ],
)
def test_truncated_refused_and_invalid_responses_not_accepted(project, provider, payload):
    with httpx.Client(
        transport=httpx.MockTransport(lambda r: httpx.Response(200, json=payload))
    ) as client:
        with pytest.raises(ProviderError):
            APIProvider(provider, "test", "key", client=client).extract(
                project.studies[0], project.fields
            )
