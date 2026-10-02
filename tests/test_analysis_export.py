import csv
import hashlib
import io
import json
import math
import warnings
import xml.etree.ElementTree as ET
import zipfile

import pytest
from matplotlib.ft2font import FT2Font
from openpyxl import load_workbook
from PIL import Image

from envevidence.analysis_data import (
    AnalysisProject,
    AnalysisStore,
    FitRequest,
    Observation,
    PlotConfig,
    ProcessingConfig,
    Series,
    import_mobile,
    process_series,
)
from envevidence.analysis_export import (
    FONT_PATH,
    ExportError,
    bundle_bytes,
    export_bundle,
    legend_mapping,
    raster_dimensions,
    render_plot,
    replay_config,
    restore_replay_sources,
    table_csv,
)
from envevidence.analysis_fitting import fit_series


@pytest.fixture
def analysis(tmp_path):
    store = AnalysisStore(tmp_path / "data")
    series = Series(name="取样 A", observations=[
        Observation(x=x, y=10 * math.exp(-0.3 * x), sigma=0.25, sample_id=f"S-{x}")
        for x in range(6)
    ])
    project = AnalysisProject(name="中文实验", series=[series], plot=PlotConfig(
        formats=["png", "tiff", "svg"], x_label="时间 (min)", y_label="浓度 (mg/L)",
    ))
    source = store.add_source(project, "原始检测.csv", b"sample,area\r\nS-001,12345\r\n")
    series.source_ids = [source.id]
    for observation in series.observations:
        observation.source_id = source.id
    project.fit_requests = [FitRequest(series_id=series.id, model="first_order")]
    result = fit_series(process_series(series), "first_order", method="linearized")
    result.update(series_id=series.id, series_name=series.name)
    project.fit_results = [result]
    project.history = [{"action": "manual_correction", "reason": "空白校正"}]
    store.save(project)
    return project, store


def test_physical_dimensions_dpi_and_chinese_are_portable(analysis):
    project, _ = analysis
    assert raster_dimensions(project.plot) == (1063, 768)
    font = FT2Font(str(FONT_PATH))
    assert all(font.get_char_index(ord(letter)) for letter in "时间浓度取样实验")
    with warnings.catch_warnings(record=True) as caught:
        for fmt in ("png", "tiff"):
            image = Image.open(io.BytesIO(render_plot(project, output_format=fmt)))
            assert image.size == (1063, 768)
            assert image.info["dpi"][0] == pytest.approx(300, abs=0.02)
        svg = render_plot(project, output_format="svg")
    assert not any("Glyph" in str(warning.message) for warning in caught)
    root = ET.fromstring(svg)
    assert float(root.attrib["width"].removesuffix("pt")) / 72 * 25.4 == pytest.approx(90)
    assert float(root.attrib["height"].removesuffix("pt")) / 72 * 25.4 == pytest.approx(65)
    assert "时间" in svg.decode("utf-8")
    assert "NotoSansCJKsc-Regular" in svg.decode("utf-8")
    assert b"<text" not in svg  # Glyph paths do not require the viewer to install the font.


def test_error_bars_and_saved_styles_change_actual_plot(analysis):
    project, _ = analysis
    original = render_plot(project)
    project.plot.error_bars = False
    assert render_plot(project) != original
    project.plot.legend = False
    project.plot.marker = "s"
    project.plot.line_style = "--"
    project.plot.x_min, project.plot.x_max = 0, 6
    assert render_plot(project) != original
    assert render_plot(project, kind="residual")
    assert render_plot(project, kind="transformed")


def test_preallocation_limit_and_missing_font_have_actionable_errors(analysis, monkeypatch):
    project, _ = analysis
    with pytest.raises(ExportError, match="50 million"):
        raster_dimensions({"width_mm": 300, "height_mm": 300, "dpi": 9600})
    with pytest.raises(ExportError, match="positive"):
        raster_dimensions({"width_mm": float("inf")})
    monkeypatch.setattr("envevidence.analysis_export.FONT_PATH", FONT_PATH.with_name("missing.otf"))
    with pytest.raises(ExportError, match="font is missing"):
        render_plot(project)


def test_bundle_contains_hashes_audit_tables_and_replayable_snapshot(analysis, tmp_path):
    project, store = analysis
    payload = bundle_bytes(project, store)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        required = {"analysis_config.json", "manifest.json", "manual_snapshot.json",
                    "project_snapshot.json", "analysis.xlsx", "processed.csv", "parameters.csv",
                    "predictions.csv", "comparison.csv", "plots/curve.png", "plots/residual.svg",
                    "plots/transformed.tiff", "curves.csv", "SOP.md"}
        assert required <= set(archive.namelist())
        manifest = json.loads(archive.read("manifest.json"))
        for name, recorded in manifest["files"].items():
            content = archive.read(name)
            assert len(content) == recorded["bytes"]
            assert hashlib.sha256(content).hexdigest() == recorded["sha256"]
        config = json.loads(archive.read("analysis_config.json"))
        source = config["project"]["sources"][0]
        assert source["relative_path"].startswith("originals/")
        assert archive.read(source["relative_path"]) == store.read_source(project, project.sources[0])
        manual = json.loads(archive.read("manual_snapshot.json"))
        assert manual["history"] == project.history
        comparison = list(csv.DictReader(io.StringIO(
            archive.read("comparison.csv").decode("utf-8-sig")
        )))
        basis = json.loads(comparison[0]["uncertainty_basis"])
        assert basis["covariance_response_scale"] == "transformed"
        assert basis["fixed_parameter_uncertainty_propagated"] is False
        sop = archive.read("SOP.md").decode("utf-8")
        assert "envevidence analyze --config analysis_config.json" in sop
        assert "not propagated" in sop
        extracted = tmp_path / "extracted"
        archive.extractall(extracted)
    restored, original_payloads = replay_config(extracted / "analysis_config.json")
    assert original_payloads[project.sources[0].id] == store.read_source(project, project.sources[0])
    assert restored.fit_requests == project.fit_requests
    assert restored.plot == project.plot
    assert restored.history == project.history
    recalculated = fit_series(process_series(restored.series[0]), "first_order", method="linearized")
    assert recalculated["parameters"]["k1"]["value"] == pytest.approx(
        project.fit_results[0]["parameters"]["k1"]["value"], rel=1e-12
    )


def test_replay_rejects_changed_originals_and_paths_outside_bundle(analysis, tmp_path):
    project, store = analysis
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        archive.extractall(tmp_path / "bundle")
    config_path = tmp_path / "bundle" / "analysis_config.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    original = config_path.parent / config["project"]["sources"][0]["relative_path"]
    original.write_bytes(b"changed")
    with pytest.raises(ExportError, match="SHA-256 mismatch"):
        replay_config(config_path)
    config["project"]["sources"][0]["relative_path"] = "../outside.csv"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    with pytest.raises(ExportError, match="relative"):
        replay_config(config_path)


def test_spreadsheet_text_is_escaped_and_numeric_values_remain_numeric(analysis):
    project, store = analysis
    project.series[0].name = "=HYPERLINK(\"danger\")"
    project.fit_results[0]["series_name"] = project.series[0].name
    rows = list(csv.DictReader(io.StringIO(table_csv([
        {"name": "  =1+1", "negative": -1.5}
    ]).decode("utf-8-sig"))))
    assert rows[0]["name"].startswith("'")
    assert rows[0]["negative"] == "-1.5"
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        workbook = load_workbook(io.BytesIO(archive.read("analysis.xlsx")))
    processed = workbook["Processed"]
    columns = {cell.value: cell.column for cell in processed[1]}
    name_cell = processed.cell(2, columns["series_name"])
    assert name_cell.value.startswith("'") and name_cell.data_type == "s"
    assert processed.cell(2, columns["y"]).data_type == "n"


def test_export_never_overwrites_and_mixed_unit_plots_are_grouped(analysis, tmp_path):
    project, store = analysis
    project.plot.formats = ["png"]
    other = Series(name="Normalized", y_unit="mg/L", observations=[
        Observation(x=0, y=5), Observation(x=1, y=4)
    ], processing=ProcessingConfig(normalize_reference=5))
    project.series.append(other)
    project.fit_results = []
    with pytest.raises(ExportError, match="same processed"):
        render_plot(project)
    first = export_bundle(project, store, tmp_path / "outputs")
    before = first.read_bytes()
    second = export_bundle(project, store, tmp_path / "outputs")
    assert first != second and first.read_bytes() == before
    with zipfile.ZipFile(first) as archive:
        assert "plots/group-1-curve.png" in archive.namelist()
        assert "plots/group-2-curve.png" in archive.namelist()


def test_independent_runs_and_transformations_retain_separate_identities(analysis):
    project, store = analysis
    project.plot.formats = ["svg"]
    result = project.fit_results[0]
    result["run_id"] = "reactor-A"
    result["request_id"] = "0:reactor-A"
    other = fit_series(process_series(project.series[0]), "second_order", method="linearized")
    other.update(series_id=project.series[0].id, series_name=project.series[0].name,
                 run_id="reactor-B", request_id="1:reactor-B")
    project.fit_results.append(other)
    with pytest.raises(ExportError, match="Different transformations"):
        render_plot(project, kind="transformed")
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        assert "plots/transformed-1.svg" in archive.namelist()
        assert "plots/transformed-2.svg" in archive.namelist()
        comparison = list(csv.DictReader(io.StringIO(
            archive.read("comparison.csv").decode("utf-8-sig")
        )))
        assert {row["run_id"] for row in comparison} == {"reactor-A", "reactor-B"}
        assert "S1/R1" in archive.read("plots/curve.svg").decode("utf-8")
        assert "S1/R2" in archive.read("plots/curve.svg").decode("utf-8")
        assert "reactor-A" in archive.read("legend_mapping.csv").decode("utf-8-sig")
        assert "reactor-B" in archive.read("legend_mapping.csv").decode("utf-8-sig")


def test_mobile_photos_and_history_are_available_after_replay(tmp_path):
    photo_name = "photos/" + "e" * 32 + ".jpg"
    photo = b"unaltered camera original"
    workspace = {"lab": {"version": 2, "experiments": [{"id": "a" * 32, "title": "Run"}],
        "samples": [{"id": "b" * 32, "experiment_id": "a" * 32, "elapsed_ms": 0,
                     "c_over_c0": 1, "revisions": [{"reason": "Measurement corrected"}]}],
        "events": [{"kind": "sample_taken"}],
        "records": [{"body": "Colour changed", "photos": [{"id": "e" * 32, "ext": "jpg"}]}]}}
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as archive:
        archive.writestr("workspace.json", json.dumps(workspace))
        archive.writestr(photo_name, photo)
    store = AnalysisStore(tmp_path / "original")
    project = AnalysisProject(name="Mobile replay")
    import_mobile(project, store, "phone.zip", raw.getvalue())
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        archive.extractall(tmp_path / "exported")
    replayed, payloads = replay_config(tmp_path / "exported" / "analysis_config.json")
    target_store = AnalysisStore(tmp_path / "replayed")
    restore_replay_sources(replayed, payloads, target_store)
    target_store.save(replayed)
    assert target_store.read_source(replayed, replayed.sources[0]) == raw.getvalue()
    member_path = replayed.mobile_archives[0]["members"][photo_name]
    assert (target_store.project_dir(replayed.id) / member_path).read_bytes() == photo
    assert replayed.mobile_archives[0]["workspace"] == workspace
    restored_source = target_store.project_dir(replayed.id) / replayed.sources[0].relative_path
    restored_source.write_bytes(b"different original")
    with pytest.raises(ExportError, match="overwrite"):
        restore_replay_sources(replayed, payloads, target_store)


def test_accepted_figures_do_not_erase_unaccepted_fit_diagnostics(analysis):
    project, store = analysis
    project.plot.formats = ["svg"]
    project.fit_results[0]["accepted"] = False
    accepted = fit_series(process_series(project.series[0]), "second_order")
    accepted.update(series_id=project.series[0].id, series_name=project.series[0].name, accepted=True)
    project.fit_results.append(accepted)
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        svg = archive.read("plots/curve.svg").decode("utf-8")
        assert "Second order" in svg and "First order" not in svg
        comparison = archive.read("comparison.csv").decode("utf-8-sig")
        assert "first_order" in comparison and "second_order" in comparison


def test_technical_sd_is_visible_and_different_observables_are_not_overlaid(tmp_path):
    series = Series(name="Technical replicates", processing=ProcessingConfig(aggregate_technical=True),
        observations=[Observation(x=x, y=y, replicate_kind="technical", sample_id=f"rep-{i}",
                                  metadata={"technical_group": f"group-{x}"})
                      for i, (x, y) in enumerate([(0, 2), (0, 4), (1, 1), (1, 3)])])
    project = AnalysisProject(name="Technical", series=[series], plot=PlotConfig(formats=["svg"]))
    svg = render_plot(project, output_format="svg").decode("utf-8")
    assert "Technical mean" in svg and "SD" in svg
    store = AnalysisStore(tmp_path)
    project.series = [Series(name="TOC", observable="TOC", observations=[Observation(x=0, y=1)]),
                      Series(name="COD", observable="COD", observations=[Observation(x=0, y=2)])]
    with pytest.raises(ExportError, match="observable"):
        render_plot(project)
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        assert "plots/group-1-curve.svg" in archive.namelist()
        assert "plots/group-2-curve.svg" in archive.namelist()


def test_export_paths_are_portable_and_cannot_escape_archive(analysis):
    project, store = analysis
    project.plot.formats = ["png"]
    project.sources[0].name = "CON.csv"
    with zipfile.ZipFile(io.BytesIO(bundle_bytes(project, store))) as archive:
        original = next(name for name in archive.namelist() if name.startswith("originals/"))
        assert original.endswith("/_CON.csv")
    project.sources[0].id = "../../unsafe"
    with pytest.raises(ExportError, match="source IDs"):
        bundle_bytes(project, store)


def test_long_default_names_do_not_fill_or_clip_default_figure(analysis):
    project, _ = analysis
    project.series[0].name = "Model compound · synthetic reaction run with a very long full title"
    project.fit_results[0].update(series_name=project.series[0].name, run_id="synthetic-run-1")
    other = fit_series(process_series(project.series[0]), "second_order")
    other.update(series_id=project.series[0].id, series_name=project.series[0].name,
                 run_id="synthetic-run-1")
    project.fit_results.append(other)
    svg = render_plot(project, output_format="svg").decode("utf-8")
    assert "Measurements" in svg and "First order" in svg and "Second order" in svg
    assert "synthetic-run-1" not in svg and "very long full title" not in svg
    mapping = legend_mapping(project)
    assert mapping[0]["series_name"] == project.series[0].name
    assert mapping[0]["run_id"] == "synthetic-run-1"
    project.series[0].conditions["plot_label"] = "自定义实验名称" * 20
    assert render_plot(project, output_format="svg")
    assert legend_mapping(project)[0]["custom_alias"] == "自定义实验名称" * 20
