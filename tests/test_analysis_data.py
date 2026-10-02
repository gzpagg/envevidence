import hashlib
import io
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from envevidence.analysis_data import (
    AnalysisProject,
    AnalysisStore,
    FitRequest,
    Observation,
    PlotConfig,
    ProcessingConfig,
    Series,
    add_manual_series,
    attach_measurements,
    demo_project,
    import_mobile,
    import_table,
    invalidate_fits,
    mobile_photos,
    parse_table,
    process_series,
    processed_units,
    remove_observation,
    revise_observation,
    series_from_rows,
    unit_factor,
)


def zip_bytes(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return output.getvalue()


def mobile_workspace():
    return {"lab": {"version": 2, "experiments": [
        {"id": "a" * 32, "title": "UV experiment", "run": {"process": "UV/H2O2"},
         "water": {"matrix": "mbr_effluent"}},
    ], "samples": [
        {"id": "b" * 32, "experiment_id": "a" * 32, "elapsed_ms": 59123,
         "recorded_elapsed_ms": 60123, "c_over_c0": 0.9,
         "ph": 7.813, "ph_source": "measured", "volume_ml": 0.015,
         "fit_excluded": False, "archived": False,
         "revisions": [{"elapsed_ms": 60123, "reason": "Corrected written time"}]},
        {"id": "c" * 32, "experiment_id": "a" * 32, "elapsed_ms": 120000,
         "c_over_c0": None, "fit_excluded": False, "archived": False},
    ], "records": [{"id": "d" * 32, "body": "Changed colour", "photos": [
        {"id": "e" * 32, "ext": "jpg"},
    ]}], "events": [{"id": "f" * 32, "kind": "sample_taken"}]}}


def test_store_preserves_raw_bytes_and_separates_literature(tmp_path):
    store = AnalysisStore(tmp_path)
    project = AnalysisProject(name="Experiment")
    data = b"\xef\xbb\xbfTime,Value\r\n0,2.0000123\r\n"
    source = store.add_source(project, "../input.csv", data)
    store.save(project)
    restored = store.load(project.id)
    assert restored == project
    assert source.name == "input.csv"
    assert source.sha256 == hashlib.sha256(data).hexdigest()
    assert store.read_source(project, source) == data
    assert store.add_source(project, "renamed.csv", data).id == source.id
    assert len(project.sources) == 1
    assert store.list_projects()[0]["id"] == project.id
    assert store.path(project.id).is_relative_to(tmp_path / "analysis")
    with pytest.raises(ValueError, match="Invalid"):
        store.project_dir("../../outside")
    target = store.project_dir(project.id) / source.relative_path
    target.write_bytes(b"changed")
    with pytest.raises(ValueError, match="hash mismatch"):
        store.read_source(project, source)


def test_csv_excel_preview_and_column_mapping(tmp_path):
    data = b"time;value;sample;run;flag\n0;4;a;rep1;valid\n1;3;b;rep1;valid\n1;2;c;rep2;valid\n"
    rows = parse_table(data, "input.csv")
    assert rows[0]["time"] == "0"
    book = Workbook()
    sheet = book.active
    sheet.append(["time", "value"])
    sheet.append([0, 3.1256789])
    output = io.BytesIO()
    book.save(output)
    assert parse_table(output.getvalue(), "input.xlsx")[0]["value"] == 3.1256789
    mapping = {"x": "time", "y": "value", "sample_id": "sample", "run_id": "run", "flag": "flag"}
    project, store = AnalysisProject(name="Runs"), AnalysisStore(tmp_path)
    series = import_table(project, store, "input.csv", data, mapping)
    assert len(series) == 2
    assert [len(s.observations) for s in series] == [2, 1]
    assert series[0].observations[0].sample_id == "a"
    assert series[1].observations[0].run_id == "rep2"
    assert import_table(project, store, "duplicate.csv", data, mapping) == []
    assert len(project.series) == 2


def test_one_original_supports_multiple_indicators_and_worksheets(tmp_path):
    project, store = AnalysisProject(name="Indicators"), AnalysisStore(tmp_path)
    data = b"time,compound,TOC\n0,10,5\n1,8,4\n2,6,3\n"
    concentration = import_table(project, store, "values.csv", data,
                                 {"x": "time", "y": "compound"})[0]
    toc = import_table(project, store, "values.csv", data,
                       {"x": "time", "y": "TOC"}, observable="TOC", y_unit="mg C/L")[0]
    assert concentration.source_ids == toc.source_ids
    assert len(project.sources) == 1
    assert toc.observations[0].y == 5
    assert import_table(project, store, "renamed.csv", data,
                        {"y": "TOC", "x": "time"}, observable="TOC", y_unit="mg C/L") == []
    book = Workbook()
    book.active.title = "Run A"
    book.active.append(["time", "value"])
    book.active.append([0, 10])
    page = book.create_sheet("Run B")
    page.append(["time", "value"])
    page.append([0, 20])
    output = io.BytesIO()
    book.save(output)
    a = import_table(project, store, "multi.xlsx", output.getvalue(),
                     {"x": "time", "y": "value"}, sheet="Run A")[0]
    b = import_table(project, store, "multi.xlsx", output.getvalue(),
                     {"x": "time", "y": "value"}, sheet="Run B")[0]
    assert a.source_ids == b.source_ids
    assert b.observations[0].y == 20
    assert len(project.sources) == 2


def test_same_manual_snapshot_can_use_distinct_scientific_templates(tmp_path):
    project, store = AnalysisProject(name="Manual"), AnalysisStore(tmp_path)
    rows, mapping = [{"x": 0, "y": 10}], {"x": "x", "y": "y"}
    first = add_manual_series(project, store, rows, mapping)[0]
    adsorption = add_manual_series(project, store, rows, mapping,
                                   kind="adsorption", observable="qt", y_unit="mg/g")[0]
    assert first.source_ids == adsorption.source_ids
    assert len(project.sources) == 1
    assert add_manual_series(project, store, rows, mapping) == []


def test_preview_rejects_formulas_duplicates_and_nonfinite_values():
    book = Workbook()
    sheet = book.active
    sheet.append(["time", "value"])
    sheet.append([0, "=1+2"])
    output = io.BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="Formula"):
        parse_table(output.getvalue(), "input.xlsx")
    with pytest.raises(ValueError, match="unique"):
        parse_table(b"x,x\n1,2", "input.csv")
    with pytest.raises(ValueError, match="finite"):
        series_from_rows([{"x": 0, "y": "inf"}], {"x": "x", "y": "y"})
    with pytest.raises(ValueError, match="mapped column"):
        series_from_rows([{"x": 0}], {"x": "x", "y": "y"})
    with pytest.raises(ValueError, match="boolean"):
        series_from_rows([{"x": True, "y": 2}], {"x": "x", "y": "y"})


def test_manual_snapshot_and_revision_history_do_not_change_original(tmp_path):
    store = AnalysisStore(tmp_path)
    project = AnalysisProject(name="Manual")
    rows = [{"time": 0, "value": 10}, {"time": 1, "value": 8}]
    series = add_manual_series(project, store, rows, {"x": "time", "y": "value"})[0]
    source = project.sources[0]
    original = store.read_source(project, source)
    project.fit_results = [{"series_id": series.id, "model": "first_order", "parameters": {"k": 0.2}}]
    project.fit_requests = [FitRequest(series_id=series.id, model="first_order", accepted=True)]
    with pytest.raises(ValueError, match="reason"):
        revise_observation(project, series.id, series.observations[0].id, {"y": 9}, " ")
    revised = revise_observation(project, series.id, series.observations[0].id,
                                  {"y": 9.123456789}, "Corrected transcription")
    assert revised.y == 9.123456789
    assert store.read_source(project, source) == original
    revision = next(h for h in project.history if h["action"] == "revise_observation")
    assert revision["before"]["y"] == 10
    assert project.fit_results == []
    assert not project.fit_requests[0].accepted
    assert any(h["action"] == "invalidate_fits" for h in project.history)
    store.save(project)
    assert store.load(project.id).series[0].observations[0].y == revised.y
    remove_observation(project, series.id, revised.id, "Duplicate transcription")
    assert len(series.observations) == 1
    assert store.read_source(project, source) == original


def test_explicit_processing_stages_units_timezero_and_sigma():
    series = Series(name="TOC", observable="TOC", x_unit="s", y_unit="ug C/L",
                    observations=[Observation(x=120, y=3000, sigma=100)], processing=ProcessingConfig(
                        blank=1000, dilution_factor=2, target_x_unit="min",
                        target_y_unit="mg C/L", time_zero=60, normalize_reference=8,
                    ))
    row = process_series(series)[0]
    assert row["original_x"] == 120
    assert row["original_y"] == 3000
    assert row["x"] == 1
    assert row["y"] == 0.5
    assert row["sigma"] == pytest.approx(0.025)
    assert row["stages"]["blank_corrected"] == 2000
    assert row["stages"]["diluted"] == 4000
    assert row["stages"]["unit_converted"] == 4
    assert processed_units(series) == ("min", "1")
    assert series.observations[0].y == 3000
    with pytest.raises(ValueError, match="Unsupported"):
        unit_factor("mg/L", "mmol/L")
    with pytest.raises(ValueError, match="Unsupported"):
        unit_factor("mg C/L", "mg O2/L")
    assert unit_factor("1/h", "1/min") == pytest.approx(1 / 60)
    assert unit_factor("mg/L/h", "ug/L/min") == pytest.approx(1000 / 60)
    assert unit_factor("mg/g/h", "mg/g/min") == pytest.approx(1 / 60)
    with pytest.raises(ValueError, match="Unsupported"):
        unit_factor("mg C/L/h", "mg O2/L/h")


def test_missing_censored_and_excluded_rows_are_not_imputed():
    rows = [{"t": 0, "y": "", "flag": "missing"},
            {"t": 1, "y": "<0.2", "flag": "below_lod"},
            {"t": 2, "y": 0.3, "flag": "below_loq"}]
    series = series_from_rows(rows, {"x": "t", "y": "y", "flag": "flag"})[0]
    processed = process_series(series)
    assert [r["included"] for r in processed] == [False, False, False]
    assert [r["y"] for r in processed] == [None, None, 0.3]
    with pytest.raises(ValidationError, match="reason"):
        Observation(x=0, y=1, excluded=True)
    excluded = Observation(x=0, y=1, excluded=True, exclusion_reason="Dropped vial")
    series.observations.append(excluded)
    assert process_series(series)[-1]["exclusion_reason"] == "Dropped vial"
    series.processing.time_zero = 1
    valid = Observation(x=0, y=1)
    series.observations.append(valid)
    assert process_series(series)[-1]["exclusion_reason"] == "before_analysis_time_zero"


def test_relative_reference_requires_actual_units_and_preserves_original():
    series = Series(name="Relative", y_unit="1", observations=[Observation(x=1, y=0.8)],
                    processing=ProcessingConfig(relative_reference=5, relative_reference_unit="mg/L"))
    assert process_series(series)[0]["y"] == 4
    assert processed_units(series)[1] == "mg/L"
    assert series.observations[0].y == 0.8
    with pytest.raises(ValidationError, match="reference requires"):
        ProcessingConfig(relative_reference=5)
    series.y_unit = "mg/L"
    with pytest.raises(ValueError, match="only used"):
        process_series(series)


def test_qt_mass_balance_correct_sigma_and_not_sampling_volume():
    with pytest.raises(ValidationError, match="mass balance"):
        ProcessingConfig(compute_qt=True)
    series = Series(name="Adsorption", kind="adsorption", observations=[
        Observation(x=2, y=4000, sigma=100, metadata={"sample_volume_ml": 1}),
    ], y_unit="ug/L", processing=ProcessingConfig(
        compute_qt=True, mass_balance_confirmed=True, adsorption_c0=10,
        reactor_volume_l=0.2, adsorbent_mass_g=0.1,
    ))
    row = process_series(series)[0]
    assert row["y"] == 12  # (10 - 4) * 0.2 / 0.1
    assert row["sigma"] == pytest.approx(0.2)
    assert row["y_unit"] == "mg/g"
    series.kind = "decay"
    with pytest.raises(ValueError, match="adsorption series"):
        process_series(series)


def test_technical_statistics_never_merge_independent_runs():
    series = Series(name="Replicates", observations=[
        Observation(x=1, y=2, sample_id="sample", run_id="run-a", replicate_kind="technical"),
        Observation(x=1, y=4, sample_id="sample", run_id="run-a", replicate_kind="technical"),
        Observation(x=1, y=20, sample_id="sample", run_id="run-b", replicate_kind="technical"),
        Observation(x=1, y=10, sample_id="separate", run_id="run-a"),
    ], processing=ProcessingConfig(aggregate_technical=True))
    rows = process_series(series)
    assert len(rows) == 3
    averaged = next(r for r in rows if r.get("technical_n"))
    assert averaged["y"] == 3
    assert averaged["technical_sd"] == pytest.approx(2 ** 0.5)
    assert averaged["sigma"] is None
    assert [r["y"] for r in rows if r["run_id"] == "run-b"] == [20]


def test_mobile_zip_preserves_corrected_times_photos_events_and_full_revisions(tmp_path):
    workspace = mobile_workspace()
    photo = b"unaltered camera bytes"
    data = zip_bytes({"workspace.json": json.dumps(workspace), f"photos/{'e' * 32}.jpg": photo})
    store, project = AnalysisStore(tmp_path), AnalysisProject(name="Mobile")
    series = import_mobile(project, store, "envbench.zip", data)
    assert len(series) == 1
    assert series[0].observations[0].x == pytest.approx(59123 / 60000)
    assert series[0].observations[0].sample_id == "b" * 32
    assert series[0].observations[0].metadata["mobile_sample"]["recorded_elapsed_ms"] == 60123
    assert series[0].observations[0].metadata["mobile_sample"]["revisions"][0]["elapsed_ms"] == 60123
    assert series[0].y_unit == "1"
    archive = project.mobile_archives[0]
    assert archive["workspace"] == workspace
    photo_path = archive["members"][f"photos/{'e' * 32}.jpg"]
    assert (store.project_dir(project.id) / photo_path).read_bytes() == photo
    assert store.read_source(project, project.sources[0]) == data
    assert import_mobile(project, store, "envbench-again.zip", data) == []
    workspace["lab"]["samples"][0]["c_over_c0"] = 0.7
    updated = zip_bytes({"workspace.json": json.dumps(workspace)})
    newer = import_mobile(project, store, "revised.zip", updated)[0]
    assert newer.id != series[0].id
    assert series[0].observations[0].y == 0.9
    assert newer.observations[0].y == 0.7
    photo_entry = mobile_photos(project, series[0].id)[0]
    assert photo_entry["member"] == f"photos/{'e' * 32}.jpg"
    assert store.read_mobile_member(project, photo_entry["source_id"], photo_entry["member"]) == photo


def test_mobile_archive_rejects_paths_and_duplicate_members_before_writing(tmp_path):
    store, project = AnalysisStore(tmp_path), AnalysisProject(name="Mobile")
    for bad_path in ("../outside", "/absolute", "photos/../../outside", "photos\\bad.jpg", "extra.txt"):
        data = zip_bytes({"workspace.json": json.dumps(mobile_workspace()), bad_path: "bad"})
        with pytest.raises(ValueError):
            import_mobile(project, store, "bad.zip", data)
    assert project.sources == []
    assert project.series == []
    assert not store.project_dir(project.id).exists()
    duplicate = mobile_workspace()
    duplicate["lab"]["experiments"].append(dict(duplicate["lab"]["experiments"][0]))
    with pytest.raises(ValueError, match="duplicate mobile experiment"):
        import_mobile(project, store, "bad.zip", zip_bytes({"workspace.json": json.dumps(duplicate)}))
    with pytest.raises(ValueError, match="workspace schema"):
        import_mobile(project, store, "bad.zip", zip_bytes({"workspace.json": "[]"}))


def test_mobile_csv_and_legacy_exclusion_without_inventing_original_reason(tmp_path):
    data = ("schema,run_id,run_title,sample_id,pulled_s,recorded_pulled_s,c_over_c0,"
            "fit_excluded,fit_exclusion_reason,volume_ml\n"
            "envbench-samples-v1,run,Experiment,sample,59.1,60.1,0.8,true,,0.015\n").encode()
    project, store = AnalysisProject(name="Mobile"), AnalysisStore(tmp_path)
    series = import_mobile(project, store, "samples.csv", data)[0]
    observation = series.observations[0]
    assert observation.x == pytest.approx(59.1 / 60)
    assert observation.excluded
    assert "not recorded" in observation.exclusion_reason
    assert observation.metadata["mobile_sample"]["fit_exclusion_reason"] == ""
    assert observation.metadata["mobile_sample"]["volume_ml"] == "0.015"


def test_attach_assays_by_mobile_sample_without_mixing_units(tmp_path):
    project, store = AnalysisProject(name="Mobile"), AnalysisStore(tmp_path)
    data = zip_bytes({"workspace.json": json.dumps(mobile_workspace())})
    series = import_mobile(project, store, "workspace.zip", data)[0]
    series.processing = ProcessingConfig(blank=0.1, dilution_factor=2, normalize_reference=0.8,
                                         relative_reference=10, relative_reference_unit="mg/L",
                                         target_y_unit="ug/L", time_zero=0.2, target_x_unit="s",
                                         aggregate_technical=True)
    rows = [{"id": "b" * 32, "TOC": 4.12345}, {"id": "c" * 32, "TOC": 3.5}]
    assert attach_measurements(project, series.id, rows, {"sample_id": "id", "y": "TOC"},
                               y_unit="mg C/L", observable="TOC", reason="Linked TOC results") == 2
    assert series.y_unit == "mg C/L"
    assert series.observable == "TOC"
    assert series.observations[0].y == 4.12345
    assert series.observations[0].metadata["mobile_sample"]["c_over_c0"] == 0.9
    assert series.processing.blank == 0
    assert series.processing.dilution_factor == 1
    assert series.processing.normalize_reference is None
    assert series.processing.relative_reference is None
    assert series.processing.target_y_unit == ""
    assert series.processing.time_zero == 0.2
    assert series.processing.target_x_unit == "s"
    assert series.processing.aggregate_technical
    assert any(h["action"] == "reset_response_processing" for h in project.history)
    with pytest.raises(ValueError, match="Unknown"):
        attach_measurements(project, series.id, [{"id": "unknown", "TOC": 1}],
                            {"sample_id": "id", "y": "TOC"}, y_unit="mg C/L",
                            observable="TOC", reason="Link")
    with pytest.raises(ValueError, match="every previously"):
        attach_measurements(project, series.id, rows[:1], {"sample_id": "id", "y": "TOC"},
                            y_unit="ug C/L", observable="TOC", reason="Link")
    assert series.observations[0].y == 4.12345


def test_plot_limits_font_settings_and_sop_templates_are_data_free(tmp_path):
    assert PlotConfig().dpi == 300
    with pytest.raises(ValidationError, match="50 million"):
        PlotConfig(width_mm=500, height_mm=500, dpi=1200)
    with pytest.raises(ValidationError, match="minimum"):
        PlotConfig(x_min=10, x_max=1)
    project = demo_project()
    store = AnalysisStore(tmp_path)
    template = store.save_template("Routine", project.model_dump())
    assert template["version"] == 1
    assert "observations" not in template["payload"]["series"][0]
    assert "conditions" not in template["payload"]["series"][0]
    assert "history" not in template["payload"]
    assert store.templates()[0] == template
    assert project.series[0].observations
    store.save(project)
    assert store.load(project.id).plot == project.plot
    assert len(store.list()) == 1


def test_invalidations_keep_old_results_as_audit_records():
    project = demo_project()
    series_id = project.series[0].id
    project.fit_results = [{"series_id": series_id, "rmse": 0.1},
                           {"series_id": "another", "rmse": 0.2}]
    invalidate_fits(project, series_id, "Processing changed")
    assert project.fit_results == [{"series_id": "another", "rmse": 0.2}]
    assert project.history[-1]["previous_results"] == [{"series_id": series_id, "rmse": 0.1}]


def test_monod_axis_cannot_apply_reaction_time_zero_or_nonfinite_processing():
    series = Series(name="Biological", kind="monod", observable="growth_rate", x_unit="mg/L",
                    y_unit="1/h", observations=[Observation(x=4, y=0.2)],
                    processing=ProcessingConfig(time_zero=1))
    with pytest.raises(ValueError, match="substrate concentration"):
        process_series(series)
    series = Series(name="Overflow", observations=[Observation(x=1, y=1e300)],
                    processing=ProcessingConfig(dilution_factor=1e100))
    with pytest.raises(ValueError, match="nonfinite"):
        process_series(series)
    with pytest.raises(ValidationError, match="hexadecimal"):
        PlotConfig(colors=["not a color"])
    with pytest.raises(ValidationError, match="nonempty"):
        Series(name=" ")


def test_actual_android_generated_csv_and_zip_compatibility(tmp_path):
    """Use the shipping phone domain APIs rather than duplicating their schema."""
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node is needed only for the cross-language mobile fixture")
    module = Path(__file__).parents[1] / "android/app/src/main/assets/lab-domain.js"
    script = r"""
const L = require(process.argv[1]);
globalThis.crypto = require('node:crypto').webcrypto;
const clock = n => ({wall:1700000000000+n, mono:10000+n, boot:'fixture'});
const lab = L.empty();
const first = L.experiment(lab, 'UV/H2O2 · 合成示例', 'Synthetic data', 'A', clock(0));
L.setRun(lab, first, {process:'UV/H2O2',target:'Model compound',oxidant:'H2O2',oxidant_mm:1}, clock(0));
L.setWater(lab, first, {matrix:'mbr_effluent',lot:'synthetic',doc_mg_c_l:6.1,ph:7.5}, clock(0));
function sample(e, ms, ratio) {
  const s = L.sample(lab,e.id,{pulled:clock(ms),quench:{agent:'MeOH',at:clock(ms+5000)},ph:7.813,ph_source:'measured',volume_ml:0.015,volume_ml_source:'measured'},clock(ms+5000));
  if (ratio !== null) L.editSample(s,{c_over_c0:ratio,peak_area:ratio*100},clock(300000));
  return s;
}
const corrected = sample(first,60123,0.876543);
L.correctSampleTimes(lab,corrected,{elapsed_ms:59123,quench_delay_ms:4000,time_revision_reason:'Aligned to written sampling record'},clock(300000));
const excluded = sample(first,120000,0.5);
L.editSample(excluded,{fit_excluded:true,fit_exclusion_reason:'Vial dropped'},clock(300000));
const empty = sample(first,180000,null);
const zero = sample(first,240000,0);
const r = L.record(lab, first.id, 'Colour changed', corrected.id, clock(260000));
L.editRecord(r,'Light brown colour',corrected.id,clock(270000));
const photoId = '9'.repeat(32);
r.photos.push({id:photoId,name:'synthetic.jpg',ext:'jpg',mime:'image/jpeg',bytes:21});
const second = L.experiment(lab,'Independent run','','B',clock(0));
sample(second,60000,0.7);
L.validate(lab);
process.stdout.write(JSON.stringify({workspace:{lab},csv:L.samplesCsv(lab),firstId:first.id,
 correctedId:corrected.id,excludedId:excluded.id,emptyId:empty.id,zeroId:zero.id,photoId}));
"""
    result = subprocess.run([node, "-e", script, str(module)], check=True, capture_output=True,
                            text=True, encoding="utf-8")
    fixture = json.loads(result.stdout)
    project, store = AnalysisProject(name="Mobile compatibility"), AnalysisStore(tmp_path)
    csv_series = import_mobile(project, store, "envbench-samples.csv", fixture["csv"].encode("utf-8"))
    first_csv = next(s for s in csv_series if s.observations[0].run_id == fixture["firstId"])
    by_id = {o.sample_id: o for o in first_csv.observations}
    assert by_id[fixture["correctedId"]].x == pytest.approx(59.1 / 60)
    assert by_id[fixture["correctedId"]].y == 0.876543
    assert by_id[fixture["excludedId"]].excluded
    assert by_id[fixture["excludedId"]].exclusion_reason == "Vial dropped"
    assert by_id[fixture["emptyId"]].y is None
    assert by_id[fixture["zeroId"]].y == 0
    assert first_csv.y_unit == "1"
    assert first_csv.conditions["process"] == "UV/H2O2"
    assert by_id[fixture["correctedId"]].metadata["mobile_sample"]["volume_ml"] == "0.015"
    assert by_id[fixture["correctedId"]].metadata["mobile_sample"]["recorded_pulled_s"] == "60.1"
    photo = b"synthetic photo bytes"
    bundle = zip_bytes({"workspace.json": json.dumps(fixture["workspace"]),
                        f"photos/{fixture['photoId']}.jpg": photo})
    zip_series = import_mobile(project, store, "envbench.zip", bundle)
    first_zip = next(s for s in zip_series if s.observations[0].run_id == fixture["firstId"])
    sample = next(o for o in first_zip.observations if o.sample_id == fixture["correctedId"])
    assert sample.x == pytest.approx(59123 / 60000)
    assert sample.metadata["mobile_sample"]["recorded_elapsed_ms"] == 60123
    assert sample.metadata["mobile_sample"]["revisions"][-1]["elapsed_ms"] == 60123
    assert project.mobile_archives[0]["workspace"]["lab"]["records"][0]["revisions"]
    assert project.mobile_archives[0]["workspace"]["lab"]["events"]
    photo_info = mobile_photos(project, first_zip.id)[0]
    assert store.read_mobile_member(project, photo_info["source_id"], photo_info["member"]) == photo
    assert import_mobile(project, store, "reimport.zip", bundle) == []
