"""Local experiment workflow; display translations never alter research inputs."""

import streamlit as st

from .analysis_data import (
    AnalysisProject,
    FitRequest,
    PlotConfig,
    ProcessingConfig,
    add_manual_series,
    attach_measurements,
    demo_project,
    import_mobile,
    import_table,
    parse_table,
    process_series,
    remove_observation,
    revise_observation,
)
from .analysis_export import export_bundle, legend_mapping, render_plot
from .analysis_fitting import run_project_fits
from .i18n import fixed_format, t
from .models import now

KINDS = {"decay": ("Degradation / removal", "降解／指标去除"),
         "adsorption": ("Adsorption", "吸附"), "monod": ("Biological rate", "生物反应速率")}
OBSERVABLES = {
    "decay": {"concentration": ("Compound concentration", "模型化合物浓度"),
              "TOC": ("TOC", "TOC"), "COD": ("COD", "COD")},
    "adsorption": {"qt": ("Adsorption capacity qt", "吸附量 qt"),
                   "concentration": ("Concentration → qt", "浓度 → 吸附量")},
    "monod": {"growth_rate": ("Specific growth rate", "比生长速率"),
              "specific_uptake_rate": ("Specific substrate uptake rate", "比底物利用速率"),
              "volumetric_rate": ("Volumetric removal rate", "体积去除速率")},
}
MODELS = {
    "zero_order": ("Zero order", "零级"), "first_order": ("First order", "一级"),
    "second_order": ("Integrated second order", "积分二级"),
    "plateau_first": ("First order with plateau", "带平台一级"),
    "adsorption_pfo": ("Adsorption pseudo-first order", "吸附拟一级"),
    "adsorption_pso": ("Adsorption pseudo-second order", "吸附拟二级"),
    "intraparticle": ("Intraparticle diffusion", "颗粒内扩散"), "monod": ("Monod", "Monod"),
}
MODEL_KINDS = {"decay": list(MODELS)[:4],
               "adsorption": ["adsorption_pfo", "adsorption_pso", "intraparticle"],
               "monod": ["monod"]}
LINEAR = {"first_order", "second_order", "adsorption_pfo", "adsorption_pso"}


def available_templates(store):
    entries = []
    for kind, indicators in OBSERVABLES.items():
        for observable in indicators:
            if kind == "adsorption" and observable == "concentration":
                continue
            unit = {"TOC": "mgC/L", "COD": "mgO2/L", "qt": "mg/g", "growth_rate": "1/h",
                    "specific_uptake_rate": "mg/g/h", "volumetric_rate": "mg/L/h"}.get(observable, "mg/L")
            entries.append({"kind": kind, "observable": observable,
                "x_unit": "mg/L" if kind == "monod" else "min", "y_unit": unit,
                "mapping": {"x": "x", "y": "y", "sample_id": "sample_id", "run_id": "run_id"},
                "processing": ProcessingConfig().model_dump(),
                "requests": [{"model": m, "method": "raw", "fixed": {}, "weighted": False,
                              "x_min": None, "x_max": None} for m in MODEL_KINDS[kind][:2]]})
    builtin = {"id": "builtin-v1", "name": t("Environmental kinetics · default v1", "环境动力学 · 默认 v1"),
               "payload": {"schema_version": 1, "plot": PlotConfig().model_dump(), "series": entries}}
    return store.templates() + [builtin]


def save(project, store, invalidate=False):
    if invalidate and project.fit_results:
        project.history.append({"action": "invalidate_fits", "previous_results": project.fit_results})
        project.fit_results = []
    if invalidate:
        for request in project.fit_requests:
            request.accepted = False
        st.session_state.pop(f"accepted_{project.id}", None)
    st.session_state.pop("analysis_export_path", None)
    store.save(project)
    st.rerun()


def error(exc):
    st.error(t("Please check the data or settings. Details:", "请检查数据或设置。详细原因：") + f" {exc}")


def diagnostic(issue):
    labels = {
        "unknown_model": "模型类型无法识别。", "unsupported_method": "此模型不支持所选拟合方法。",
        "invalid_weighting": "加权只适用于原始尺度拟合。", "invalid_fixed": "请检查固定参数的名称、数值和物理范围。",
        "invalid_interval": "请填写有效的拟合区间。", "missing_exclusion_reason": "排除记录需要填写理由。",
        "invalid_row": "参与拟合的自变量和响应必须是有限的非负数值。",
        "invalid_sigma": "加权拟合需要每个参与点的正测量标准差。",
        "insufficient_points": "至少需要四个有效点，且有效点数应不小于自由参数数加二。",
        "identical_x": "至少需要两个不同的自变量值。", "covariance_warning": "参数协方差估计存在问题。",
        "not_identifiable": "参数无法可靠区分，已不显示置信区间。",
        "parameter_correlation": "参数高度相关，置信区间可能无法可靠估计。",
        "parameter_at_bound": "参数位于物理边界，已不显示对称置信区间。",
        "negative_prediction": "零级模型在拟合区间内预测了负响应，请检查适用范围。",
        "constant_response": "响应值相同，原始尺度 R² 无法定义。", "fit_failed": "拟合未成功，请检查数据和模型。",
        "missing_series": "拟合所需的数据系列不存在。", "model_family_mismatch": "模型与当前实验类型不匹配。",
        "processing_failed": "数据处理未成功，请检查单位和处理规则。",
        "adsorption_requires_qt": "吸附拟合需要 qt，请先输入吸附量或确认质量平衡后由浓度计算 qt。",
        "invalid_monod_axis": "Monod 的自变量应为底物浓度，而不是时间。",
        "invalid_monod_observable": "Monod 需要比生长速率、比利用速率或体积去除速率。",
        "monod_requires_substrate_axis": "Monod 的自变量应为底物浓度，而不是时间，请明确提供浓度单位。",
        "monod_requires_rate": "Monod 需要比生长速率、比利用速率或体积去除速率。",
    }
    if isinstance(issue, dict):
        transformed = {
            "ln(C) requires every included response to be positive.": "ln(C) 要求所有参与点的响应大于零，请检查或明确排除不适用点。",
            "1/C requires every included response to be positive.": "1/C 要求所有参与点的响应大于零，请检查或明确排除不适用点。",
            "PFO linearization requires an independently supplied fixed qe.": "拟一级线性化需要独立提供固定 qe，不能使用最后一个测量点代替。",
            "ln(qe-qt) requires every included qt to be less than fixed qe.": "ln(qe−qt) 要求参与点的 qt 小于固定 qe，请检查参考值或明确排除不适用点。",
            "t/qt requires every included time and qt to be positive.": "t/qt 要求参与点的时间和 qt 均大于零，请明确处理 t=0 或 qt=0 的记录。",
        }
        if issue.get("message") in transformed:
            return t(issue["message"], transformed[issue["message"]])
        return t(issue.get("message", ""), labels.get(issue.get("code"), issue.get("message", "")))
    return str(issue)


def column(label_en, label_zh, key, columns, required=False, mapping=None):
    choices = columns if required else [""] + columns
    saved = (mapping or {}).get(key)
    guess = choices.index(saved) if saved in choices else next((i for i, item in enumerate(choices) if item.lower() == key), 0)
    return st.selectbox(t(label_en, label_zh), choices, index=guess,
                        key=f"map_{key}", format_func=fixed_format(lambda x: x or t("None", "无")))


def input_spec(prefix):
    cols = st.columns(3)
    kind = cols[0].selectbox(t("Experiment type", "实验类型"), list(KINDS),
                            format_func=fixed_format(lambda k: t(*KINDS[k])), key=f"{prefix}_kind")
    observable = cols[1].selectbox(t("Measured indicator", "测量指标"), list(OBSERVABLES[kind]),
                                  format_func=fixed_format(lambda k: t(*OBSERVABLES[kind][k])),
                                  key=f"{prefix}_observable_{kind}")
    unit_default = {"TOC": "mgC/L", "COD": "mgO2/L", "qt": "mg/g",
                    "growth_rate": "1/h", "specific_uptake_rate": "mg/g/h",
                    "volumetric_rate": "mg/L/h"}.get(observable, "mg/L")
    y_unit = cols[2].text_input(t("Response unit", "响应单位"), value=unit_default,
                                key=f"{prefix}_unit_{observable}")
    cols = st.columns(2)
    if kind == "monod":
        x_unit = cols[0].text_input(t("Substrate concentration unit", "底物浓度单位"), "mg/L",
                                    key=f"{prefix}_xunit_monod")
    else:
        x_unit = cols[0].selectbox(t("Time unit", "时间单位"), ["min", "s", "h"],
                                   key=f"{prefix}_xunit_time")
    series_name = cols[1].text_input(t("New series name", "新系列名称"),
                                     value=t("Treatment A", "处理 A"), key=f"{prefix}_name")
    return dict(series_name=series_name, kind=kind, observable=observable,
                x_unit=x_unit, y_unit=y_unit)


def imports(project, store):
    st.subheader(t("Add a measurement series", "添加测量数据系列"))
    mode = st.radio(t("Input source", "输入来源"), ["manual", "file", "mobile"], horizontal=True,
                    format_func=fixed_format(lambda x: t(*{
                        "manual": ("Enter / paste", "录入／粘贴"),
                        "file": ("CSV / Excel", "CSV／Excel"),
                        "mobile": ("EnvBench CSV / ZIP", "EnvBench CSV／ZIP")}[x])), key="analysis_input")
    if mode == "mobile":
        file = st.file_uploader(t("Phone experiment export", "手机实验导出文件"),
                                type=["csv", "zip"], key="mobile_upload")
        st.caption(t("ZIP keeps photos, events and revisions. Ratios remain relative until a reference is provided.",
                     "ZIP 保留照片、事件和修订记录。C/C0 保持相对值，提供参考浓度后才可转换。"))
        if st.button(t("Import phone archive", "导入手机档案"), disabled=file is None,
                     key="analysis_import_mobile"):
            try:
                import_mobile(project, store, file.name, file.getvalue())
                save(project, store, True)
            except (ValueError, OSError) as exc:
                error(exc)
        return
    spec = input_spec("input")
    chosen_sop = st.selectbox(t("Import SOP", "导入 SOP"), [None] + available_templates(store),
        format_func=fixed_format(lambda s: s["name"] if s else t("Manual settings", "手动配置")), key="import_sop")
    sop_config = next((s for s in chosen_sop["payload"]["series"]
        if all(s[k] == spec[k] for k in ("kind", "observable", "x_unit", "y_unit"))), None) if chosen_sop else None
    saved_mapping = sop_config.get("mapping", {}) if sop_config else {}
    if chosen_sop and not sop_config:
        st.warning(t("This SOP does not match the selected input indicator and units. Choose manual settings or a matching SOP.",
                     "此 SOP 与所选指标或输入单位不匹配，请选择手动配置或匹配的 SOP。"))
    if mode == "manual":
        text = st.text_area(t("Paste CSV or tab-separated measurements", "粘贴 CSV 或制表符分隔数据"),
                            "x,y,sample_id,run_id\n0,10,S0,R1\n5,7.4,S1,R1\n10,5.5,S2,R1\n15,4.1,S3,R1\n20,3.0,S4,R1\n30,1.65,S5,R1",
                            height=175, key="analysis_paste")
        data, filename = text.encode("utf-8"), "manual.csv"
    else:
        file = st.file_uploader(t("Quantified measurements", "已定量的测量结果"),
                                type=["csv", "xlsx"], key="analysis_upload")
        if file is None:
            return
        data, filename = file.getvalue(), file.name
    try:
        sheet = None
        if filename.lower().endswith(".xlsx"):
            from io import BytesIO

            from openpyxl import load_workbook

            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
            sheet = st.selectbox(t("Worksheet", "工作表"), book.sheetnames, key="analysis_sheet")
            book.close()
        rows = parse_table(data, filename, sheet)
        if not rows:
            st.info(t("Add a header and measurement rows.", "请添加表头和测量数据行。"))
            return
        st.dataframe(rows[:12], hide_index=True, width="stretch")
        columns = list(rows[0])
        map_context = (project.id, chosen_sop["id"] if chosen_sop else None, spec["kind"],
                       spec["observable"], spec["x_unit"], spec["y_unit"], filename, sheet, tuple(columns))
        if st.session_state.get("analysis_map_context") != map_context:
            for logical in ("x", "y", "sample_id", "run_id", "sigma", "flag", "excluded", "exclusion_reason", "replicate_kind"):
                st.session_state.pop(f"map_{logical}", None)
                if saved_mapping.get(logical) in columns:
                    st.session_state[f"map_{logical}"] = saved_mapping[logical]
            st.session_state.analysis_map_context = map_context
        with st.expander(t("Column mapping", "字段映射"), expanded=True):
            left, right = st.columns(2)
            with left:
                x = column("Time / substrate", "时间／底物浓度", "x", columns, True, saved_mapping)
                sample = column("Sample ID", "样品编号", "sample_id", columns, mapping=saved_mapping)
                flag = column("Detection flag", "检出状态", "flag", columns, mapping=saved_mapping)
                excluded = column("Excluded", "排除标记", "excluded", columns, mapping=saved_mapping)
            with right:
                y = column("Response", "响应数值", "y", columns, True, saved_mapping)
                run = column("Independent run ID", "独立实验编号", "run_id", columns, mapping=saved_mapping)
                sigma = column("Measurement standard deviation", "测量标准差", "sigma", columns, mapping=saved_mapping)
                reason = column("Exclusion reason", "排除理由", "exclusion_reason", columns, mapping=saved_mapping)
            replicate = column("Replicate type", "重复类型", "replicate_kind", columns, mapping=saved_mapping)
        mapping = {k: v for k, v in {"x": x, "y": y, "sample_id": sample, "run_id": run,
                   "sigma": sigma, "flag": flag, "excluded": excluded,
                   "exclusion_reason": reason, "replicate_kind": replicate}.items() if v}
        if st.button(t("Save measurement series", "保存测量数据系列"), type="primary", key="analysis_import",
                     disabled=chosen_sop is not None and sop_config is None):
            if mode == "manual":
                added = add_manual_series(project, store, rows=rows, mapping=mapping, **spec)
            else:
                added = import_table(project, store, filename, data, mapping=mapping, sheet=sheet, **spec)
            if sop_config:
                for series in added:
                    series.processing = ProcessingConfig.model_validate(sop_config["processing"])
                    project.fit_requests.extend(FitRequest(series_id=series.id, **r) for r in sop_config["requests"])
                    process_series(series)
                project.plot = PlotConfig.model_validate(chosen_sop["payload"]["plot"])
            save(project, store, True)
    except (ValueError, OSError, KeyError) as exc:
        error(exc)


def conditions_page(project, series, store):
    st.caption(t("Conditions describe this series. They do not change measured values.",
                 "反应条件属于当前系列，不会改变测量数值。"))
    fields = [("pollutant", "Compound / substrate", "化合物／底物"),
              ("process", "Treatment process", "处理工艺"), ("water_matrix", "Water matrix", "水体基质"),
              ("reagent_dose", "Reagent / catalyst dose (with unit)", "药剂／催化剂剂量（含单位）"),
              ("pH", "pH", "pH"), ("temperature", "Temperature (with unit)", "温度（含单位）"),
              ("reactor_volume", "Reactor liquid volume (with unit)", "反应液体积（含单位）"),
              ("biomass", "Biomass / dry mass (with unit)", "生物量／干质量（含单位）"),
              ("DO", "Dissolved oxygen (with unit)", "溶解氧（含单位）"),
              ("rate_origin", "Rate measurement / calculation", "速率测量／计算来源")]
    fields.append(("plot_label", "Custom plot label (optional)", "自定义图例名称（可选）"))
    with st.form(f"conditions_{series.id}"):
        name = st.text_input(t("Series name", "数据系列名称"), series.name)
        cols = st.columns(2)
        values = {}
        for i, (key, en, zh) in enumerate(fields):
            values[key] = cols[i % 2].text_input(t(en, zh), str(series.conditions.get(key, "")))
        values["notes"] = st.text_area(t("Experimental notes", "实验说明"), str(series.conditions.get("notes", "")))
        if st.form_submit_button(t("Save conditions", "保存反应条件"), type="primary"):
            if not name.strip():
                st.error(t("Enter a series name.", "请填写数据系列名称。"))
            else:
                project.history.append({"at": now(), "action": "conditions", "series_id": series.id,
                    "before": {"name": series.name, "conditions": series.conditions.copy()},
                    "after": {"name": name.strip(), "conditions": {**series.conditions, **values}}})
                series.name = name.strip()
                series.conditions.update(values)
                save(project, store, True)


def data_page(project, series, store):
    st.caption(t("Machine field names stay stable in exports. Sample IDs connect phone records and measurements.",
                 "导出的机器字段名保持稳定。样品编号用于连接手机记录与检测结果。"))
    entries = [{k: getattr(o, k) for k in ("id", "sample_id", "run_id", "x", "y", "sigma", "flag",
                                          "replicate_kind", "excluded", "exclusion_reason")}
               for o in series.observations]
    edited = st.data_editor(entries, hide_index=True, width="stretch", disabled=["id", "sample_id", "run_id"],
                            key=f"observations_{series.id}", column_config={
                                "x": st.column_config.NumberColumn(t("Time / substrate", "时间／底物浓度")),
                                "y": st.column_config.NumberColumn(t("Measured response", "测量响应")),
                                "sigma": st.column_config.NumberColumn(t("Standard deviation", "标准差")),
                                "flag": st.column_config.SelectboxColumn(t("Detection", "检出状态"),
                                                options=["valid", "missing", "below_lod", "below_loq"],
                                                format_func=fixed_format(lambda v: t(*{"valid": ("Valid", "有效"), "missing": ("Missing", "缺失"),
                                                    "below_lod": ("Below LOD", "低于检出限"), "below_loq": ("Below LOQ", "低于定量限")}[v]))),
                                "replicate_kind": st.column_config.SelectboxColumn(t("Replicate", "重复类型"),
                                                options=["independent", "technical"],
                                                format_func=fixed_format(lambda v: t("Independent run", "独立实验") if v == "independent" else t("Technical replicate", "技术重复"))),
                                "excluded": st.column_config.CheckboxColumn(t("Exclude", "排除")),
                                "exclusion_reason": st.column_config.TextColumn(t("Exclusion reason", "排除理由")),
                            })
    reason = st.text_input(t("Reason for data correction", "数据修订理由"), key=f"revision_reason_{series.id}")
    if st.button(t("Save data corrections", "保存数据修订"), key="analysis_revision"):
        try:
            candidate = project.model_copy(deep=True)
            for original, changed in zip(entries, edited):
                changes = {k: v for k, v in changed.items() if k not in {"id", "sample_id", "run_id"} and v != original[k]}
                if changes:
                    revise_observation(candidate, series.id, original["id"], changes, reason)
            save(candidate, store, True)
        except ValueError as exc:
            error(exc)
    with st.expander(t("Remove a measurement with an audit record", "移除测量记录并保留审计历史")):
        if series.observations:
            removed = st.selectbox(t("Measurement to remove", "要移除的测量记录"), series.observations,
                                   format_func=fixed_format(lambda o: f"{o.sample_id} · {o.x}"), key="analysis_remove_select")
            why = st.text_input(t("Removal reason", "移除理由"), key="analysis_remove_reason")
            if st.button(t("Remove measurement", "移除测量记录"), key="analysis_remove"):
                try:
                    remove_observation(project, series.id, removed.id, why)
                    save(project, store, True)
                except ValueError as exc:
                    error(exc)
    with st.expander(t("Link quantified results by sample ID", "通过样品编号关联检测结果")):
        assay = st.file_uploader(t("Assay CSV / Excel", "检测结果 CSV／Excel"), type=["csv", "xlsx"], key="assay_upload")
        if assay is not None:
            try:
                assay_rows = parse_table(assay.getvalue(), assay.name)
                if assay_rows:
                    st.dataframe(assay_rows[:10], hide_index=True, width="stretch")
                    cols = list(assay_rows[0])
                    sample_col = st.selectbox(t("Sample ID column", "样品编号列"), cols, key="assay_sample")
                    value_col = st.selectbox(t("Measured value column", "检测值列"), cols, key="assay_y")
                    sigma_col = st.selectbox(t("SD column (optional)", "标准差列（可选）"), [""] + cols, key="assay_sigma")
                    flag_col = st.selectbox(t("Detection column (optional)", "检出状态列（可选）"), [""] + cols, key="assay_flag")
                    unit = st.text_input(t("Assay unit", "检测结果单位"), "mg/L", key="assay_unit")
                    observable = st.selectbox(t("Assay indicator", "检测指标"), list(OBSERVABLES[series.kind]),
                                             format_func=fixed_format(lambda k: t(*OBSERVABLES[series.kind][k])), key="assay_observable")
                    why = st.text_input(t("Reason / assay reference", "关联理由／检测来源"), key="assay_reason")
                    if st.button(t("Link measurements", "关联检测结果"), key="assay_attach"):
                        candidate = project.model_copy(deep=True)
                        source = store.add_source(candidate, assay.name, assay.getvalue())
                        attach_measurements(candidate, series.id, assay_rows,
                            {"sample_id": sample_col, "y": value_col, "sigma": sigma_col, "flag": flag_col},
                            y_unit=unit, observable=observable, reason=why, source_id=source.id)
                        save(candidate, store, True)
            except (ValueError, OSError) as exc:
                error(exc)
    with st.expander(t("Imported phone archive", "已导入的手机档案")):
        archives = [a for a in project.mobile_archives if a["source_id"] in series.source_ids]
        st.json({"archives": [a["workspace"] for a in archives],
                 "sample_metadata": [o.metadata for o in series.observations if o.metadata]})
        photos = [(a["source_id"], name) for a in archives for name in a["members"]
                  if name.startswith("photos/") and ".thumb." not in name]
        if photos:
            from io import BytesIO
            from zipfile import ZipFile

            source_id, member = st.selectbox(t("Photo", "照片"), photos,
                                  format_func=fixed_format(lambda p: p[1]), key="analysis_photo")
            source = next(s for s in project.sources if s.id == source_id)
            with ZipFile(BytesIO(store.read_source(project, source))) as archive:
                st.image(archive.read(member), width=500)


def processing_page(project, series, store):
    p = series.processing
    with st.form(f"processing_{series.id}"):
        cols = st.columns(3)
        blank = cols[0].number_input(t("Blank (measured units)", "空白值（测量单位）"), value=float(p.blank))
        dilution = cols[1].number_input(t("Dilution factor", "稀释倍数"), min_value=0.000001, value=float(p.dilution_factor))
        zero = cols[2].number_input(t("Time zero offset (input time unit)", "时间零点偏移（输入时间单位）"),
                                   value=float(p.time_zero), disabled=series.kind == "monod")
        cols = st.columns(2)
        xu = cols[0].text_input(t("Processed x unit (empty = unchanged)", "处理后自变量单位（空白则保留）"), p.target_x_unit)
        yu = cols[1].text_input(t("Processed response unit (empty = unchanged)", "处理后响应单位（空白则保留）"), p.target_y_unit)
        normalized = st.checkbox(t("Normalize by an explicit reference", "使用明确参考值归一化"), value=p.normalize_reference is not None)
        ref = st.number_input(t("Normalization reference (processed unit)", "归一化参考值（处理后单位）"), min_value=0.000001,
                              value=float(p.normalize_reference or 1))
        relative = st.checkbox(t("Convert C/C0 using an absolute reference", "使用绝对参考浓度转换 C/C0"), value=p.relative_reference is not None)
        cols = st.columns(2)
        absolute = cols[0].number_input(t("Absolute reference concentration", "绝对参考浓度"), min_value=0.000001,
                                       value=float(p.relative_reference or 1))
        absolute_unit = cols[1].text_input(t("Absolute reference unit", "绝对参考浓度单位"), p.relative_reference_unit or "mg/L")
        aggregate = st.checkbox(t("Aggregate technical replicates of the same sample", "汇总同一样品的技术重复"), value=p.aggregate_technical)
        qt = p.compute_qt
        c0, volume, mass, confirmed = p.adsorption_c0, p.reactor_volume_l, p.adsorbent_mass_g, p.mass_balance_confirmed
        if series.kind == "adsorption":
            qt = st.checkbox(t("Calculate qt from concentrations", "通过浓度计算 qt"), value=qt)
            cols = st.columns(3)
            c0 = cols[0].number_input(t("Initial concentration (mg/L)", "初始浓度（mg/L）"), min_value=0.0, value=float(c0 or 10))
            volume = cols[1].number_input(t("Reactor liquid volume (L)", "反应液体积（L）"), min_value=0.000001, value=float(volume or 0.1))
            mass = cols[2].number_input(t("Adsorbent dry mass (g)", "吸附剂干质量（g）"), min_value=0.000001, value=float(mass or 0.1))
            confirmed = st.checkbox(t("Constant-volume / separate-bottle mass balance applies; no other removal process", "适用恒体积／独立瓶质量平衡，且没有其他去除过程"), value=confirmed)
        st.caption(t("Correction: (measurement − blank) × dilution. Defaults leave measurements unchanged.",
                     "校正顺序：（测量值 − 空白）× 稀释倍数。默认保留原始测量值。"))
        if st.form_submit_button(t("Save processing rules", "保存处理规则"), type="primary"):
            try:
                draft = ProcessingConfig(blank=blank, dilution_factor=dilution, time_zero=zero,
                    target_x_unit=xu, target_y_unit=yu, normalize_reference=ref if normalized else None,
                    relative_reference=absolute if relative else None,
                    relative_reference_unit=absolute_unit if relative else "", aggregate_technical=aggregate,
                    compute_qt=qt, adsorption_c0=c0, reactor_volume_l=volume,
                    adsorbent_mass_g=mass, mass_balance_confirmed=confirmed)
                check = series.model_copy(update={"processing": draft})
                process_series(check)
                project.history.append({"at": now(), "action": "processing_rules", "series_id": series.id,
                    "before": p.model_dump(), "after": draft.model_dump()})
                series.processing = draft
                save(project, store, True)
            except ValueError as exc:
                error(exc)
    try:
        st.subheader(t("Processed data preview", "处理后数据预览"))
        st.dataframe(process_series(series), width="stretch", hide_index=True)
    except ValueError as exc:
        error(exc)


def fitting_page(project, series, store):
    choices = MODEL_KINDS[series.kind]
    prior = [r for r in project.fit_requests if r.series_id == series.id]
    previous = [r.model for r in prior]
    selected = st.multiselect(t("Candidate models", "候选模型"), choices,
                              default=list(dict.fromkeys(previous)) or choices[:2],
                              format_func=fixed_format(lambda k: t(*MODELS[k])), key=f"models_{series.id}")
    saved_methods = {r.method for r in prior}
    default_method = "both" if len(saved_methods) == 2 else (next(iter(saved_methods)) if saved_methods else "raw")
    method = st.selectbox(t("Fitting method", "拟合方法"), ["raw", "both", "linearized"],
                          index=["raw", "both", "linearized"].index(default_method),
                          format_func=fixed_format(lambda k: t(*{"raw": ("Original response scale", "原始响应尺度"),
                              "both": ("Original + applicable linearizations", "原始尺度＋适用线性化"),
                              "linearized": ("Applicable linearizations", "适用线性化")}[k])), key=f"method_{series.id}")
    weighted = st.checkbox(t("Weight raw fits by provided measurement SD", "原始尺度按已提供测量标准差加权"),
                           value=any(r.weighted for r in prior), key=f"weighted_{series.id}")
    fixed = {}
    if series.kind in {"decay", "adsorption"}:
        key = "C0" if series.kind == "decay" else "qe"
        previous_fixed = next((r.fixed[key] for r in prior if key in r.fixed), None)
        if st.checkbox(t(f"Fix {key} to an independently measured value", f"将 {key} 固定为独立测量值"),
                       value=previous_fixed is not None, key=f"fix_{series.id}"):
            fixed[key] = st.number_input(t(f"Fixed {key} (processed response unit)", f"固定 {key}（处理后响应单位）"),
                                         min_value=0.000001, value=float(previous_fixed or 10), key=f"fixed_value_{series.id}")
    interval = st.checkbox(t("Select fitting interval", "指定拟合区间"), key=f"interval_{series.id}")
    lower = upper = None
    if interval:
        cols = st.columns(2)
        lower = cols[0].number_input(t("From (processed x unit)", "起点（处理后自变量单位）"), value=0.0)
        upper = cols[1].number_input(t("To (processed x unit)", "终点（处理后自变量单位）"), value=30.0)
    st.caption(t("Compare errors on the original scale. Linearized R² is shown separately. Independent runs are fitted separately.",
                 "比较采用原始尺度误差；线性化 R² 单独显示。独立实验分别拟合。"))
    if st.button(t("Fit selected models", "拟合所选模型"), type="primary", key="analysis_fit", disabled=not selected):
        requests = []
        for model in selected:
            methods = ["raw"] if method == "raw" else (["raw", "linearized"] if method == "both" else ["linearized"])
            for m in methods:
                if m == "linearized" and model not in LINEAR:
                    continue
                requests.append(FitRequest(series_id=series.id, model=model, method=m,
                    fixed={} if model == "intraparticle" else fixed,
                    weighted=weighted if m == "raw" else False, x_min=lower, x_max=upper))
        try:
            if not requests:
                st.warning(t("Selected models use the original scale. Choose that method to fit them.",
                             "所选模型使用原始尺度，请选择相应方法进行拟合。"))
            else:
                if project.fit_results:
                    project.history.append({"action": "refit", "previous_results": project.fit_results})
                st.session_state.pop(f"accepted_{project.id}", None)
                project.fit_requests = [r for r in project.fit_requests if r.series_id != series.id] + requests
                run_project_fits(project)
                save(project, store)
        except ValueError as exc:
            error(exc)
    results = [r for r in project.fit_results if r.get("series_id") == series.id]
    if not results:
        return
    comparison = []
    for result in results:
        comparison.append({"run": result.get("run_id"), "model": t(*MODELS.get(result.get("model"), ("", ""))),
                           "method": t("Original scale", "原始尺度") if result.get("method") == "raw" else t("Linearized", "线性化"),
                           "status": t(*{"ok": ("Complete", "完成"), "warning": ("Review diagnostics", "需检查诊断"),
                                        "failed": ("Failed", "失败")}.get(result.get("status"), ("", ""))),
                           **result.get("metrics", {})})
    st.dataframe(comparison, hide_index=True, width="stretch", column_config={
        "run": t("Run", "实验编号"), "model": t("Model", "模型"), "method": t("Method", "方法"),
        "status": t("Status", "状态"), "n_points": t("Points", "有效点数"), "n_parameters": t("Free parameters", "自由参数数"),
        "rmse": st.column_config.NumberColumn("RMSE", format="%.4g"),
        "mae": st.column_config.NumberColumn("MAE", format="%.4g"),
        "r2": st.column_config.NumberColumn("R²", format="%.5g"),
    })
    for result in results:
        method_label = t("Original scale", "原始尺度") if result.get("method") == "raw" else t("Linearized", "线性化")
        label = f"{result.get('run_id', '')} · {t(*MODELS.get(result.get('model'), ('', '')))} · {method_label}"
        with st.expander(label):
            for issue in result.get("errors", []):
                st.error(diagnostic(issue))
            for issue in result.get("warnings", []):
                st.warning(diagnostic(issue))
            params = result.get("parameters", {})
            st.dataframe([{"parameter": k, **(v if isinstance(v, dict) else {"value": v})}
                          for k, v in params.items()] if isinstance(params, dict) else params,
                         width="stretch", hide_index=True, column_config={
                             "parameter": t("Parameter", "参数"), "value": t("Value", "数值"),
                             "unit": t("Unit", "单位"), "ci95": t("95% interval", "95% 区间"), "fixed": t("Fixed", "固定")})
            if result.get("transformed_metrics"):
                st.write(t("Transformed-scale metrics", "变换尺度指标"), result["transformed_metrics"])


def charts_page(project, store):
    p = project.plot
    with st.form("plot_config"):
        cols = st.columns(4)
        fmt = cols[0].selectbox(t("Format", "格式"), ["png", "tiff", "svg"], index=["png", "tiff", "svg"].index(p.format))
        width = cols[1].number_input(t("Width (mm)", "宽度（mm）"), min_value=10.0, value=float(p.width_mm))
        height = cols[2].number_input(t("Height (mm)", "高度（mm）"), min_value=10.0, value=float(p.height_mm))
        dpi = cols[3].number_input("DPI", min_value=72, max_value=9600, value=int(p.dpi), step=100,
                                   help=t("Typical presets: 300 / 600 / 1200", "常用值：300／600／1200"))
        with st.expander(t("Advanced plot style", "高级绘图格式")):
            cols = st.columns(2)
            xlabel = cols[0].text_input(t("X axis label (empty = automatic)", "X 轴标题（空白则自动）"), p.x_label)
            ylabel = cols[1].text_input(t("Y axis label (empty = automatic)", "Y 轴标题（空白则自动）"), p.y_label)
            cols = st.columns(4)
            font_size = cols[0].number_input(t("Font size (pt)", "字号（pt）"), min_value=4.0, value=float(p.font_size))
            line_width = cols[1].number_input(t("Line width (pt)", "线宽（pt）"), min_value=0.1, value=float(p.line_width))
            marker = cols[2].selectbox(t("Marker", "点形"), ["o", "s", "^", "D", "x", "+"], index=["o", "s", "^", "D", "x", "+"].index(p.marker))
            marker_size = cols[3].number_input(t("Marker size (pt)", "点大小（pt）"), min_value=1.0, value=float(p.marker_size))
            font = st.selectbox(t("Font", "字体"), ["Noto Sans CJK SC", "DejaVu Sans"], index=0 if p.font == "Noto Sans CJK SC" else 1)
            palette = st.text_input(t("Curve colors (comma-separated hex)", "曲线颜色（逗号分隔十六进制）"), ",".join(p.colors))
            cols = st.columns(2)
            legend = cols[0].checkbox(t("Show legend", "显示图例"), p.legend)
            bars = cols[1].checkbox(t("Show measurement error bars", "显示测量误差条"), p.error_bars)
            limits = st.checkbox(t("Set axis ranges", "设置坐标轴范围"), value=any(v is not None for v in [p.x_min, p.x_max, p.y_min, p.y_max]))
            cols = st.columns(4)
            bounds = [cols[i].number_input(label, value=float(value or 0)) for i, (label, value) in enumerate([
                (t("X minimum", "X 最小值"), p.x_min), (t("X maximum", "X 最大值"), p.x_max),
                (t("Y minimum", "Y 最小值"), p.y_min), (t("Y maximum", "Y 最大值"), p.y_max)])]
        if st.form_submit_button(t("Save plot style", "保存绘图格式"), type="primary"):
            try:
                project.plot = PlotConfig(format=fmt, formats=[fmt], width_mm=width, height_mm=height, dpi=dpi,
                    font=font, font_size=font_size, line_width=line_width, marker=marker,
                    marker_size=marker_size, colors=[s.strip() for s in palette.split(",") if s.strip()],
                    x_label=xlabel, y_label=ylabel, legend=legend, error_bars=bars,
                    **dict(zip(["x_min", "x_max", "y_min", "y_max"], bounds if limits else [None] * 4)))
                save(project, store)
            except ValueError as exc:
                error(exc)
    if project.fit_results:
        kind = st.selectbox(t("Plot", "图像类型"), ["curve", "residual", "transformed"],
                            format_func=fixed_format(lambda k: t(*{"curve": ("Measurements and curves", "数据点与拟合曲线"),
                                "residual": ("Residuals", "残差"), "transformed": ("Linearized fits", "线性化拟合")}[k])), key="analysis_plot_kind")
        selected_ids = st.multiselect(t("Series to plot together", "叠图数据系列"), [s.id for s in project.series],
                                       default=[project.series[0].id],
                                       format_func=fixed_format(lambda sid: next(s.name for s in project.series if s.id == sid)), key="plot_series")
        try:
            preview = render_plot(project, kind=kind, output_format="png", series_ids=selected_ids)
            st.image(preview, width=650)
            with st.expander(t("Full legend identities", "完整图例对应记录")):
                st.dataframe(legend_mapping(project, series_ids=selected_ids), hide_index=True, width="stretch")
            data = render_plot(project, kind=kind, series_ids=selected_ids)
            st.download_button(t("Download this plot", "下载当前图像"), data, f"{kind}.{p.format}", key="analysis_plot_download")
        except ValueError as exc:
            error(exc)


def sop_page(project, store):
    st.subheader(t("Reusable SOP template", "可复用 SOP 模板"))
    templates = available_templates(store)
    cols = st.columns(2)
    template_name = cols[0].text_input(t("Template name", "模板名称"), "Water-treatment kinetics", key="analysis_template_name")
    if cols[0].button(t("Save SOP template", "保存 SOP 模板"), key="analysis_template_save"):
        payload = {"schema_version": 1, "plot": project.plot.model_dump(),
                   "series": [{"kind": s.kind, "observable": s.observable, "x_unit": s.x_unit,
                       "y_unit": s.y_unit, "processing": s.processing.model_dump(),
                       "mapping": s.conditions.get("column_mapping", {}),
                       "requests": [r.model_dump(exclude={"series_id", "accepted"}) for r in project.fit_requests if r.series_id == s.id]}
                      for s in project.series]}
        store.save_template(template_name, payload)
        st.session_state.pop("analysis_template_select", None)
        st.rerun()
    if templates:
        selected = cols[1].selectbox(t("Saved SOP", "已保存 SOP"), templates,
                                     format_func=fixed_format(lambda s: s["name"]), key="analysis_template_select")
        if cols[1].button(t("Apply SOP to matching series", "将 SOP 应用于匹配系列"), key="analysis_template_apply"):
            payload = selected["payload"]
            candidate = project.model_copy(deep=True)
            candidate.plot = PlotConfig.model_validate(payload["plot"])
            matches = 0
            for series in candidate.series:
                cfg = next((c for c in payload["series"] if c["kind"] == series.kind and c["observable"] == series.observable
                            and c["x_unit"] == series.x_unit and c["y_unit"] == series.y_unit), None)
                if cfg:
                    matches += 1
                    series.processing = ProcessingConfig.model_validate(cfg["processing"])
                    series.conditions["column_mapping"] = cfg.get("mapping", {})
                    candidate.fit_requests = [r for r in candidate.fit_requests if r.series_id != series.id]
                    candidate.fit_requests.extend(FitRequest(series_id=series.id, **r) for r in cfg["requests"])
                    process_series(series)
            if matches:
                candidate.history.append({"at": now(), "action": "apply_sop", "template_id": selected["id"],
                    "before": {"plot": project.plot.model_dump(), "processing": {s.id: s.processing.model_dump() for s in project.series}},
                    "after": {"plot": candidate.plot.model_dump(), "processing": {s.id: s.processing.model_dump() for s in candidate.series}}})
                save(candidate, store, True)
            else:
                st.warning(t("No series matches this SOP's experiment type, indicator and input units.",
                             "没有与此 SOP 的实验类型、指标和输入单位相匹配的系列。"))
    st.divider()
    st.subheader(t("Confirm and export analysis", "确认并导出分析成果"))
    st.write(t("The package contains original files, processed tables, parameters, predictions, residuals, figures and a replayable configuration.",
               "成果包包含原始文件、处理数据、参数、预测、残差、图表及可重跑的分析配置。"))
    results = project.fit_results
    options = [i for i, r in enumerate(results) if r.get("status") in {"warning", "ok"}]
    accepted = st.multiselect(t("Results to include in accepted figures", "确认用于成果图的拟合结果"), options,
                             default=[i for i in options if results[i].get("accepted")],
                             format_func=fixed_format(lambda i: f"{results[i].get('series_name', '')} / {results[i].get('run_id', '')} / {results[i].get('model')} / {results[i].get('method')}"), key=f"accepted_{project.id}")
    confirm = st.checkbox(t("I checked the conditions, processing rules and selected fits", "已检查反应条件、处理规则与所选拟合结果"), key="analysis_confirm")
    if st.button(t("Create SOP package", "生成 SOP 成果包"), type="primary", disabled=not confirm or not accepted, key="analysis_export"):
        try:
            for i, result in enumerate(results):
                result["accepted"] = i in accepted
            store.save(project)
            path = export_bundle(project, store, store.root / "exports")
            st.session_state.analysis_export_path = str(path)
        except (ValueError, OSError) as exc:
            error(exc)
    path = st.session_state.get("analysis_export_path")
    if path:
        from pathlib import Path

        path = Path(path)
        if path.exists():
            st.download_button(t("Download SOP package", "下载 SOP 成果包"), path.read_bytes(), path.name,
                               mime="application/zip", key="analysis_bundle_download")
    with st.expander(t("Audit and configuration", "审计记录与配置")):
        st.json({"sources": [s.model_dump() for s in project.sources], "history": project.history})


def analysis_page(store):
    st.header(t("Experiment analysis", "实验分析"))
    st.markdown("<div class='ee-hero'><h3 style='color:inherit'>" + t("From reaction conditions to reproducible curves.", "从反应条件，到可复现的动力学曲线。") + "</h3><p>" + t("Import · Process · Fit · Review · Export", "导入 · 处理 · 拟合 · 检查 · 导出") + "</p></div>", unsafe_allow_html=True)
    cols = st.columns([1, 1, 2])
    if cols[0].button(t("New analysis project", "新建分析项目"), key="analysis_new"):
        st.session_state.pop("analysis_project_id", None)
        st.session_state.pop("analysis_export_path", None)
    if cols[1].button(t("Load analysis demo", "载入分析演示"), key="analysis_demo"):
        project = demo_project(store)
        store.save(project)
        st.session_state.analysis_project_id = project.id
        st.session_state.pop("analysis_export_path", None)
        st.rerun()
    projects = store.list_projects()
    if projects:
        selected = cols[2].selectbox(t("Saved analysis projects", "已保存的分析项目"), projects,
                                    format_func=fixed_format(lambda p: p["name"]), key="analysis_project_select")
        if cols[2].button(t("Open analysis project", "打开分析项目"), key="analysis_open"):
            st.session_state.analysis_project_id = selected["id"]
            st.session_state.pop("analysis_export_path", None)
            st.rerun()
    if not st.session_state.get("analysis_project_id"):
        with st.form("new_analysis"):
            name = st.text_input(t("Project name", "项目名称"), t("My experiment", "我的实验"))
            if st.form_submit_button(t("Create analysis project", "创建分析项目"), type="primary"):
                if name.strip():
                    project = AnalysisProject(name=name.strip())
                    store.save(project)
                    st.session_state.analysis_project_id = project.id
                    st.rerun()
        return
    try:
        project = store.load(st.session_state.analysis_project_id)
    except (ValueError, OSError) as exc:
        error(exc)
        return
    st.subheader(project.name)
    cols = st.columns(3)
    cols[0].metric(t("Series", "数据系列"), len(project.series))
    cols[1].metric(t("Measurements", "测量记录"), sum(len(s.observations) for s in project.series))
    cols[2].metric(t("Model fits", "拟合结果"), len(project.fit_results))
    with st.expander(t("Import / add data", "导入／添加数据"), expanded=not project.series):
        imports(project, store)
    if not project.series:
        return
    def series_label(sid):
        current = next(s for s in project.series if s.id == sid)
        indicator = OBSERVABLES[current.kind].get(current.observable, (current.observable, current.observable))
        return f"{current.name} · {t(*indicator)}"

    series_id = st.selectbox(t("Current series", "当前数据系列"), [s.id for s in project.series],
                          format_func=fixed_format(series_label), key="analysis_series")
    series = next(s for s in project.series if s.id == series_id)
    tabs = st.tabs([t("Conditions", "反应条件"), t("Measurements", "测量数据"), t("Processing", "数据处理"),
                    t("Fitting", "动力学拟合"), t("Charts", "图表格式"), t("SOP export", "SOP 导出")])
    functions = [conditions_page, data_page, processing_page, fitting_page]
    for tab, function in zip(tabs[:4], functions):
        with tab:
            function(project, series, store)
    with tabs[4]:
        charts_page(project, store)
    with tabs[5]:
        sop_page(project, store)
