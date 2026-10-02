"""Independent, local experiment inputs, provenance, and explicit processing.

Literature project schemas are deliberately not reused. Raw uploads are immutable;
analysis changes live in observations, processing settings, and audit events.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import statistics
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path, PurePosixPath
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from .models import now, uid

MAX_UPLOAD_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_BYTES = 150 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 256


class AnalysisModel(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)


class SourceFile(AnalysisModel):
    id: str = Field(default_factory=uid, pattern=r"^[a-f0-9]{32}$")
    name: str
    sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    relative_path: str
    kind: Literal["upload", "manual", "mobile"] = "upload"
    imported_at: str = Field(default_factory=now)


class Observation(AnalysisModel):
    id: str = Field(default_factory=uid)
    sample_id: str = Field(default_factory=uid)
    run_id: str = "run-1"
    replicate_kind: Literal["independent", "technical"] = "independent"
    x: float | None = None
    y: float | None = None
    sigma: float | None = Field(default=None, gt=0)
    flag: Literal["valid", "missing", "below_lod", "below_loq"] = "valid"
    excluded: bool = False
    exclusion_reason: str = ""
    source_id: str = ""
    metadata: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode="after")
    def reason_required(self):
        if self.excluded and not self.exclusion_reason.strip():
            raise ValueError("Excluded observations require a reason")
        return self


class ProcessingConfig(AnalysisModel):
    blank: float = 0
    dilution_factor: float = Field(default=1, gt=0)
    target_x_unit: str = ""
    target_y_unit: str = ""
    time_zero: float = 0
    normalize_reference: float | None = Field(default=None, gt=0)
    relative_reference: float | None = Field(default=None, gt=0)
    relative_reference_unit: str = ""
    compute_qt: bool = False
    adsorption_c0: float | None = Field(default=None, ge=0)
    reactor_volume_l: float | None = Field(default=None, gt=0)
    adsorbent_mass_g: float | None = Field(default=None, gt=0)
    mass_balance_confirmed: bool = False
    aggregate_technical: bool = False

    @model_validator(mode="after")
    def reference_unit_required(self):
        if self.relative_reference is not None and not self.relative_reference_unit.strip():
            raise ValueError("A relative concentration reference requires its unit")
        if self.compute_qt:
            if not self.mass_balance_confirmed:
                raise ValueError("Confirm the constant-volume adsorption mass balance")
            if any(v is None for v in (
                self.adsorption_c0, self.reactor_volume_l, self.adsorbent_mass_g,
            )):
                raise ValueError("qt calculation requires C0, reactor volume, and dry mass")
            if self.normalize_reference is not None:
                raise ValueError("qt calculation and normalization are separate workflows")
        return self


class PlotConfig(AnalysisModel):
    format: Literal["png", "tiff", "svg"] = "png"
    formats: list[Literal["png", "tiff", "svg"]] = Field(default_factory=lambda: ["png"])
    width_mm: float = Field(default=90, gt=0)
    height_mm: float = Field(default=65, gt=0)
    dpi: int = Field(default=300, ge=72, le=9600)
    font: str = "Noto Sans CJK SC"
    font_size: float = Field(default=9, gt=0)
    line_width: float = Field(default=1.4, gt=0)
    line_style: Literal["-", "--", "-.", ":"] = "-"
    marker: str = "o"
    marker_size: float = Field(default=4, gt=0)
    colors: list[str] = Field(default_factory=lambda: ["#A65338", "#147D73", "#1D4ED8"])
    x_label: str = ""
    y_label: str = ""
    x_min: float | None = None
    x_max: float | None = None
    y_min: float | None = None
    y_max: float | None = None
    legend: bool = True
    error_bars: bool = True
    show_curve: bool = True
    show_residual: bool = True
    show_transformed: bool = True

    @model_validator(mode="after")
    def limits_and_pixels(self):
        for low, high in ((self.x_min, self.x_max), (self.y_min, self.y_max)):
            if low is not None and high is not None and low >= high:
                raise ValueError("Axis minimum must be below its maximum")
        if self.width_mm / 25.4 * self.dpi * self.height_mm / 25.4 * self.dpi > 50_000_000:
            raise ValueError("Raster export exceeds 50 million pixels; reduce size or DPI")
        if not self.formats:
            raise ValueError("Select at least one export format")
        if not self.colors or any(not re.fullmatch(r"#[a-fA-F0-9]{6}", c) for c in self.colors):
            raise ValueError("Plot colors must be six-digit hexadecimal colors")
        if self.marker not in {"o", "s", "^", "D", "x", "+", ".", "v", "*"}:
            raise ValueError("Unsupported plot marker")
        return self


class FitRequest(AnalysisModel):
    series_id: str
    model: str
    method: Literal["raw", "linearized"] = "raw"
    fixed: dict[str, float] = Field(default_factory=dict)
    weighted: bool = False
    x_min: float | None = None
    x_max: float | None = None
    accepted: bool = False


class Series(AnalysisModel):
    id: str = Field(default_factory=uid)
    name: str
    kind: Literal["decay", "adsorption", "monod"] = "decay"
    observable: Literal[
        "concentration", "TOC", "COD", "qt", "growth_rate", "specific_uptake_rate",
        "volumetric_rate",
    ] = "concentration"
    x_unit: str = "min"
    y_unit: str = "mg/L"
    conditions: dict[str, Any] = Field(default_factory=dict)
    observations: list[Observation] = Field(default_factory=list)
    processing: ProcessingConfig = Field(default_factory=ProcessingConfig)
    source_ids: list[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def nonempty_labels(self):
        if not self.name.strip() or not self.x_unit.strip() or not self.y_unit.strip():
            raise ValueError("Series name and units must be nonempty")
        if len({o.id for o in self.observations}) != len(self.observations):
            raise ValueError("Observation IDs within a series must be unique")
        return self


class AnalysisProject(AnalysisModel):
    schema_version: Literal[1] = 1
    id: str = Field(default_factory=uid)
    name: str
    created_at: str = Field(default_factory=now)
    updated_at: str = Field(default_factory=now)
    series: list[Series] = Field(default_factory=list)
    sources: list[SourceFile] = Field(default_factory=list)
    fit_requests: list[FitRequest] = Field(default_factory=list)
    fit_results: list[dict[str, Any]] = Field(default_factory=list)
    plot: PlotConfig = Field(default_factory=PlotConfig)
    history: list[dict[str, Any]] = Field(default_factory=list)
    mobile_archives: list[dict[str, Any]] = Field(default_factory=list)
    sop_templates: list[dict[str, Any]] = Field(default_factory=list)

    @model_validator(mode="after")
    def nonempty_name(self):
        if not self.name.strip():
            raise ValueError("Project name must be nonempty")
        for items in (self.series, self.sources):
            if len({item.id for item in items}) != len(items):
                raise ValueError("Project item IDs must be unique")
        return self


def _atomic_write(target: Path, content: bytes) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=target.parent, prefix=".save-", suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


class AnalysisStore:
    """Atomic project JSON and content-addressed originals under data/analysis."""

    def __init__(self, root: str | Path = "data"):
        self.root = Path(root) / "analysis"
        self.root.mkdir(parents=True, exist_ok=True)

    def project_dir(self, project_id: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{32}", project_id):
            raise ValueError("Invalid analysis project id")
        return self.root / project_id

    def path(self, project_id: str) -> Path:
        return self.project_dir(project_id) / "project.json"

    def save(self, project: AnalysisProject) -> Path:
        project.updated_at = now()
        validated = AnalysisProject.model_validate(project.model_dump())
        target = self.path(project.id)
        _atomic_write(target, validated.model_dump_json(indent=2).encode("utf-8"))
        return target

    def load(self, project_id: str) -> AnalysisProject:
        return AnalysisProject.model_validate_json(self.path(project_id).read_bytes())

    def list_projects(self) -> list[dict[str, str]]:
        items = []
        for path in self.root.glob("*/project.json"):
            try:
                project = AnalysisProject.model_validate_json(path.read_bytes())
                if path == self.path(project.id):
                    items.append({
                        "id": project.id, "name": project.name, "updated_at": project.updated_at,
                    })
            except (ValueError, OSError):
                continue
        return sorted(items, key=lambda item: item["updated_at"], reverse=True)

    def list(self) -> list[dict[str, str]]:
        return self.list_projects()

    def save_template(self, name: str, payload: dict[str, Any]) -> dict[str, Any]:
        if not name.strip():
            raise ValueError("A SOP template needs a name")
        blocked = {"observations", "sources", "source_ids", "fit_results", "history",
                   "mobile_archives", "conditions", "raw", "data"}

        def settings_only(value):
            if isinstance(value, dict):
                return {k: settings_only(v) for k, v in value.items() if k not in blocked}
            if isinstance(value, list):
                return [settings_only(v) for v in value]
            return value

        template = {"id": uid(), "name": name.strip(), "version": 1,
                    "created_at": now(), "payload": settings_only(payload)}
        content = json.dumps(template, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")
        if len(content) > 1024 * 1024:
            raise ValueError("SOP template is too large")
        _atomic_write(self.root / "templates" / f"{template['id']}.json", content)
        return template

    def templates(self) -> list[dict[str, Any]]:
        result = []
        for path in (self.root / "templates").glob("*.json"):
            try:
                template = json.loads(path.read_bytes())
                if template.get("version") == 1 and template.get("id") == path.stem:
                    result.append(template)
            except (OSError, ValueError):
                continue
        return sorted(result, key=lambda item: item.get("created_at", ""), reverse=True)

    def add_source(
        self, project: AnalysisProject, name: str, data: bytes, kind: str = "upload",
    ) -> SourceFile:
        if len(data) > MAX_UPLOAD_BYTES:
            raise ValueError("Upload exceeds 100 MiB")
        digest = hashlib.sha256(data).hexdigest()
        existing = next((s for s in project.sources if s.sha256 == digest), None)
        if existing:
            return existing
        clean_name = Path(name.replace("\\", "/")).name or "input.bin"
        suffix = Path(clean_name).suffix.lower()
        if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
            suffix = ".bin"
        relative = f"raw/{digest}{suffix}"
        target = self.project_dir(project.id) / relative
        if target.exists() and target.read_bytes() != data:
            raise ValueError("Original file hash conflict")
        if not target.exists():
            _atomic_write(target, data)
        source = SourceFile(name=clean_name, sha256=digest, relative_path=relative, kind=kind)
        project.sources.append(source)
        return source

    def read_source(self, project: AnalysisProject, source: SourceFile) -> bytes:
        base = self.project_dir(project.id).resolve()
        target = (base / source.relative_path).resolve()
        if not target.is_relative_to(base):
            raise ValueError("Unsafe original file path")
        data = target.read_bytes()
        if hashlib.sha256(data).hexdigest() != source.sha256:
            raise ValueError(f"Original file hash mismatch: {source.name}")
        return data

    def read_mobile_member(
        self, project: AnalysisProject, source_id: str, member: str,
    ) -> bytes:
        archive = next(a for a in project.mobile_archives if a["source_id"] == source_id)
        if member not in archive["members"]:
            raise ValueError("Unknown mobile archive member")
        source = next(s for s in project.sources if s.id == source_id)
        # Recover from the immutable original, validating its hash, so replayed projects
        # do not depend on extracted photo copies remaining at their old local path.
        data = self.read_source(project, source)
        with zipfile.ZipFile(io.BytesIO(data)) as zipped:
            return zipped.read(member)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or _text(value).lower() in ("", "na", "n/a", "nan", "null", "none"):
        return None
    if isinstance(value, bool):
        raise ValueError("A boolean is not a numeric measurement")
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("Measurements must be finite")
    return number


def _boolean(value: Any) -> bool:
    text = _text(value).lower()
    if text in ("", "false", "0", "no"):
        return False
    if text in ("true", "1", "yes"):
        return True
    raise ValueError(f"Invalid boolean: {value}")


def parse_table(data: bytes, filename: str, sheet: str | None = None) -> list[dict[str, Any]]:
    """Preview CSV/TSV/XLSX without evaluating or trusting spreadsheet formula cells."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("Upload exceeds 100 MiB")
    if Path(filename).suffix.lower() in (".xlsx", ".xlsm"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_ARCHIVE_BYTES:
                raise ValueError("Worksheet archive expands beyond 150 MiB")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            page = book[sheet] if sheet else book.worksheets[0]
            if page.max_column and page.max_column > MAX_COLUMNS:
                raise ValueError("Table exceeds 256 columns")
            table = []
            for row in page.iter_rows():
                if len(table) > MAX_ROWS:
                    raise ValueError("Table exceeds 100000 rows")
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Formula cells require an exported values-only worksheet")
                table.append([cell.value for cell in row])
        finally:
            book.close()
    else:
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in (text.splitlines() or [""])[0] else csv.excel
        table = list(csv.reader(io.StringIO(text), dialect))
        if len(table) > MAX_ROWS + 1:
            raise ValueError("Table exceeds 100000 rows")
    while table and not any(_text(v) for v in table[-1]):
        table.pop()
    if not table:
        raise ValueError("The table is empty")
    headers = [_text(v) for v in table[0]]
    if len(headers) > MAX_COLUMNS:
        raise ValueError("Table exceeds 256 columns")
    if not headers or any(not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError("Column names must be nonempty and unique")
    rows = []
    for number, values in enumerate(table[1:], start=2):
        if not any(_text(v) for v in values):
            continue
        if len(values) > len(headers) and any(_text(v) for v in values[len(headers):]):
            raise ValueError(f"Row {number} has more cells than the header")
        rows.append(dict(zip(headers, list(values) + [None] * (len(headers) - len(values)))))
    return rows


def _flag(value: Any, y: Any) -> str:
    label = _text(value).lower().replace(" ", "_")
    aliases = {"<lod": "below_lod", "<loq": "below_loq", "nd": "below_lod"}
    label = aliases.get(label, label)
    if not label:
        label = "missing" if _number(y) is None else "valid"
    if label not in ("valid", "missing", "below_lod", "below_loq"):
        raise ValueError(f"Unknown measurement flag: {value}")
    return label


def series_from_rows(
    rows: list[dict[str, Any]], mapping: dict[str, str], *, series_name: str = "Measurements",
    kind: str = "decay", observable: str = "concentration", x_unit: str = "min",
    y_unit: str = "mg/L", source_id: str = "",
) -> list[Series]:
    """Group independent runs; preserve technical-repeat observations within each run."""
    if not mapping.get("x") or not mapping.get("y"):
        raise ValueError("Map the independent variable and measurement columns")
    groups: dict[str, list[Observation]] = defaultdict(list)
    for index, row in enumerate(rows):
        if any(column not in row for column in mapping.values() if column):
            raise ValueError("A mapped column is absent from the input")
        def get(key):
            return row.get(mapping.get(key, ""))
        run = _text(get("run_id")) or "run-1"
        sample = _text(get("sample_id")) or f"{source_id or 'manual'}-{index + 1}"
        raw_y = get("y")
        flag = _flag(get("flag"), raw_y)
        # Censored strings stay in metadata; an optional numeric bound is not an observed zero.
        y = None if flag in ("below_lod", "below_loq") and _text(raw_y).startswith("<") else _number(raw_y)
        replicate = _text(get("replicate_kind")) or "independent"
        groups[run].append(Observation(
            sample_id=sample, run_id=run, x=_number(get("x")), y=y,
            sigma=_number(get("sigma")), flag=flag, replicate_kind=replicate,
            excluded=_boolean(get("excluded")), exclusion_reason=_text(get("exclusion_reason")),
            source_id=source_id, metadata={"source_row": index + 2, "input": row},
        ))
    return [Series(
        name=f"{series_name} · {run}" if len(groups) > 1 else series_name,
        kind=kind, observable=observable, x_unit=x_unit, y_unit=y_unit,
        observations=observations, source_ids=[source_id] if source_id else [],
    ) for run, observations in groups.items()]


def import_table(
    project: AnalysisProject, store: AnalysisStore, filename: str, data: bytes,
    mapping: dict[str, str], *, series_name: str = "Measurements", kind: str = "decay",
    observable: str = "concentration", x_unit: str = "min", y_unit: str = "mg/L",
    sheet: str | None = None,
) -> list[Series]:
    digest = hashlib.sha256(data).hexdigest()
    signature = _import_signature(digest, mapping, kind, observable, x_unit, y_unit, sheet)
    if any(s.conditions.get("import_signature") == signature for s in project.series):
        return []
    rows = parse_table(data, filename, sheet)
    series = series_from_rows(rows, mapping, series_name=series_name, kind=kind,
                              observable=observable, x_unit=x_unit, y_unit=y_unit)
    source = store.add_source(project, filename, data)
    for item in series:
        item.source_ids = [source.id]
        item.conditions["column_mapping"] = dict(mapping)
        item.conditions["import_signature"] = signature
        for observation in item.observations:
            observation.source_id = source.id
            if not mapping.get("sample_id"):
                observation.sample_id = f"{source.id}-{observation.metadata['source_row'] - 1}"
    project.series.extend(series)
    project.history.append({"at": now(), "action": "import", "source_id": source.id,
                            "mapping": mapping, "series_ids": [s.id for s in series]})
    return series


def _import_signature(digest, mapping, kind, observable, x_unit, y_unit, sheet=None):
    settings = {"sha256": digest, "mapping": {k: v for k, v in mapping.items() if v},
                "kind": kind, "observable": observable, "x_unit": x_unit,
                "y_unit": y_unit, "sheet": sheet}
    return hashlib.sha256(json.dumps(settings, sort_keys=True).encode("utf-8")).hexdigest()


def add_manual_series(
    project: AnalysisProject, store: AnalysisStore, rows: list[dict[str, Any]],
    mapping: dict[str, str], **options,
) -> list[Series]:
    data = json.dumps({"rows": rows, "mapping": mapping}, ensure_ascii=False,
                      sort_keys=True, allow_nan=False).encode("utf-8")
    digest = hashlib.sha256(data).hexdigest()
    signature = _import_signature(digest, mapping, options.get("kind", "decay"),
                                  options.get("observable", "concentration"),
                                  options.get("x_unit", "min"), options.get("y_unit", "mg/L"))
    if any(s.conditions.get("import_signature") == signature for s in project.series):
        return []
    series = series_from_rows(rows, mapping, **options)
    source = store.add_source(project, "manual-input.json", data, "manual")
    for item in series:
        item.source_ids = [source.id]
        item.conditions["column_mapping"] = dict(mapping)
        item.conditions["import_signature"] = signature
        for observation in item.observations:
            observation.source_id = source.id
            if not mapping.get("sample_id"):
                observation.sample_id = f"{source.id}-{observation.metadata['source_row'] - 1}"
    project.series.extend(series)
    project.history.append({"at": now(), "action": "manual_snapshot", "source_id": source.id,
                            "mapping": mapping, "series_ids": [s.id for s in series]})
    return series


def revise_observation(
    project: AnalysisProject, series_id: str, observation_id: str,
    changes: dict[str, Any], reason: str,
) -> Observation:
    if not reason.strip():
        raise ValueError("A revision requires a reason")
    allowed = {"x", "y", "sigma", "flag", "excluded", "exclusion_reason", "replicate_kind"}
    if not set(changes).issubset(allowed):
        raise ValueError("Source identity and raw metadata cannot be edited")
    series = next(s for s in project.series if s.id == series_id)
    index = next(i for i, o in enumerate(series.observations) if o.id == observation_id)
    before = series.observations[index]
    revised = Observation.model_validate({**before.model_dump(), **changes})
    project.history.append({"at": now(), "action": "revise_observation", "reason": reason,
                            "series_id": series_id, "observation_id": observation_id,
                            "before": before.model_dump(), "after": revised.model_dump()})
    series.observations[index] = revised
    invalidate_fits(project, series_id, reason)
    return revised


def invalidate_fits(project: AnalysisProject, series_id: str, reason: str) -> None:
    previous = [r for r in project.fit_results if r.get("series_id") == series_id]
    if previous:
        project.history.append({"at": now(), "action": "invalidate_fits", "reason": reason,
                                "series_id": series_id, "previous_results": previous})
    project.fit_results = [r for r in project.fit_results if r.get("series_id") != series_id]
    for request in project.fit_requests:
        if request.series_id == series_id:
            request.accepted = False


def remove_observation(
    project: AnalysisProject, series_id: str, observation_id: str, reason: str,
) -> None:
    if not reason.strip():
        raise ValueError("Removing an observation requires a reason")
    series = next(s for s in project.series if s.id == series_id)
    before = next(o for o in series.observations if o.id == observation_id)
    project.history.append({"at": now(), "action": "remove_observation", "reason": reason,
                            "series_id": series_id, "before": before.model_dump()})
    series.observations = [o for o in series.observations if o.id != observation_id]
    invalidate_fits(project, series_id, reason)


_UNITS = {
    "s": ("time", 1 / 60), "sec": ("time", 1 / 60), "min": ("time", 1),
    "h": ("time", 60), "hr": ("time", 60), "d": ("time", 1440),
    "mg/l": ("mass_concentration", 1), "g/l": ("mass_concentration", 1000),
    "ug/l": ("mass_concentration", 0.001), "ng/l": ("mass_concentration", 0.000001),
    "mol/l": ("molar_concentration", 1000), "mmol/l": ("molar_concentration", 1),
    "umol/l": ("molar_concentration", 0.001), "m": ("molar_concentration", 1000),
    "mm": ("molar_concentration", 1), "um": ("molar_concentration", 0.001),
    "mgc/l": ("carbon_concentration", 1), "gc/l": ("carbon_concentration", 1000),
    "ugc/l": ("carbon_concentration", 0.001),
    "mgo2/l": ("oxygen_concentration", 1), "go2/l": ("oxygen_concentration", 1000),
    "ugo2/l": ("oxygen_concentration", 0.001),
    "mg/g": ("adsorption", 1), "g/g": ("adsorption", 1000),
    "1": ("ratio", 1), "ratio": ("ratio", 1), "c/c0": ("ratio", 1),
}


def _unit_key(unit: str) -> str:
    return unit.strip().lower().replace("μ", "u").replace("µ", "u").replace(" ", "").replace("₂", "2")


def unit_factor(source: str, target: str) -> float:
    source_key, target_key = _unit_key(source), _unit_key(target)
    if source_key == target_key:
        return 1.0

    def specification(key):
        if key in _UNITS:
            return _UNITS[key]
        # Explicit concentration/capacity rates and specific growth rates. No
        # conversion between mass, carbon, oxygen demand, or molar bases.
        numerator, separator, denominator = key.rpartition("/")
        base, time = _UNITS.get(numerator), _UNITS.get(denominator)
        if separator and base and time and time[0] == "time" and base[0] != "time":
            return f"rate:{base[0]}", base[1] / time[1]
        return None

    a, b = specification(source_key), specification(target_key)
    if a is None or b is None or a[0] != b[0]:
        raise ValueError(f"Unsupported unit conversion: {source} -> {target}")
    return a[1] / b[1]


def processed_units(series: Series) -> tuple[str, str]:
    config = series.processing
    x_unit = config.target_x_unit or series.x_unit
    y_unit = config.relative_reference_unit if config.relative_reference is not None else series.y_unit
    y_unit = config.target_y_unit or y_unit
    if config.compute_qt:
        y_unit = "mg/g"
    elif config.normalize_reference is not None:
        y_unit = "1"
    return x_unit, y_unit


def process_series(series: Series) -> list[dict[str, Any]]:
    """Retain all rows and every numeric stage; fitters use only included rows."""
    config = series.processing
    # Revalidate edited config models before any calculation.
    config = ProcessingConfig.model_validate(config.model_dump())
    if series.kind == "monod" and config.time_zero != 0:
        raise ValueError("A substrate concentration axis cannot use a reaction time offset")
    x_unit, final_y_unit = processed_units(series)
    xf = unit_factor(series.x_unit, x_unit)
    current_y_unit = series.y_unit
    reference_factor = 1.0
    if config.relative_reference is not None:
        if _UNITS.get(_unit_key(series.y_unit), (None,))[0] != "ratio":
            raise ValueError("An absolute reference is only used for relative concentration inputs")
        current_y_unit = config.relative_reference_unit
        reference_factor = config.relative_reference
    conversion_y_unit = config.target_y_unit or current_y_unit
    yf = unit_factor(current_y_unit, conversion_y_unit)
    qt_factor = unit_factor(conversion_y_unit, "mg/L") if config.compute_qt else 1.0
    if config.compute_qt and series.kind != "adsorption":
        raise ValueError("qt calculation belongs to an adsorption series")
    rows = []
    for observation in series.observations:
        x, y, sigma = observation.x, observation.y, observation.sigma
        stages: dict[str, Any] = {"raw_x": x, "raw_y": y, "raw_sigma": sigma}
        if x is not None:
            x = (x - config.time_zero) * xf
        if y is not None:
            y = (y - config.blank) * config.dilution_factor
            stages["blank_corrected"] = observation.y - config.blank
            stages["diluted"] = y
            y *= reference_factor
            stages["absolute_reference_applied"] = y
            y *= yf
            stages["unit_converted"] = y
            if config.compute_qt:
                factor = config.reactor_volume_l / config.adsorbent_mass_g
                y = (config.adsorption_c0 - y * qt_factor) * factor
                stages["qt"] = y
            elif config.normalize_reference is not None:
                y /= config.normalize_reference
                stages["normalized"] = y
        if sigma is not None:
            sigma *= config.dilution_factor * reference_factor * abs(yf)
            if config.compute_qt:
                sigma *= qt_factor * config.reactor_volume_l / config.adsorbent_mass_g
            elif config.normalize_reference is not None:
                sigma /= config.normalize_reference
        if any(value is not None and not math.isfinite(value) for value in (x, y, sigma)):
            raise ValueError("Processing produced a nonfinite value; check factors and units")
        reason = observation.exclusion_reason if observation.excluded else ""
        if not reason and observation.flag != "valid":
            reason = observation.flag
        if not reason and (x is None or y is None):
            reason = "missing"
        if not reason and series.kind != "monod" and x is not None and x < 0:
            reason = "before_analysis_time_zero"
        rows.append({
            "observation_id": observation.id, "sample_ids": [observation.sample_id],
            "sample_id": observation.sample_id, "run_id": observation.run_id,
            "replicate_kind": observation.replicate_kind, "source_id": observation.source_id,
            "x": x, "y": y, "sigma": sigma, "original_x": observation.x,
            "original_y": observation.y, "x_unit": x_unit, "y_unit": final_y_unit,
            "included": not bool(reason), "exclusion_reason": reason, "flag": observation.flag,
            "stages": stages, "metadata": observation.metadata,
        })
    if not config.aggregate_technical:
        return rows
    groups: dict[tuple, list[dict]] = defaultdict(list)
    output = []
    for row in rows:
        if row["included"] and row["replicate_kind"] == "technical":
            group_id = row["metadata"].get("technical_group") or row["sample_id"]
            groups[(row["run_id"], group_id, row["x"])].append(row)
        else:
            output.append(row)
    for members in groups.values():
        if len(members) == 1:
            output.extend(members)
            continue
        row = dict(members[0])
        values = [member["y"] for member in members]
        row.update({"y": statistics.mean(values), "sigma": None,
                    "sample_ids": [m["sample_id"] for m in members],
                    "technical_n": len(members), "technical_sd": statistics.stdev(values),
                    "observation_ids": [m["observation_id"] for m in members],
                    "stages": {"technical_members": members}})
        # SD describes spread, not a supplied uncertainty of the mean for weighted fitting.
        output.append(row)
    return sorted(output, key=lambda r: (r["run_id"], r["x"] is None, r["x"] or 0))


def _mobile_csv(rows: list[dict], source_id: str) -> list[Series]:
    groups: dict[str, list[Observation]] = defaultdict(list)
    conditions: dict[str, dict] = {}
    names: dict[str, str] = {}
    seen = set()
    for index, row in enumerate(rows):
        if row.get("schema") != "envbench-samples-v1":
            raise ValueError("Unsupported EnvBench sample schema")
        run, sample = _text(row.get("run_id")), _text(row.get("sample_id"))
        if not run or not sample or (run, sample) in seen:
            raise ValueError("Mobile sample identities must be nonempty and unique")
        seen.add((run, sample))
        excluded = _boolean(row.get("fit_excluded"))
        reason = _text(row.get("fit_exclusion_reason"))
        if excluded and not reason:
            reason = "Imported exclusion; original reason not recorded"
        y = _number(row.get("c_over_c0"))
        seconds = _number(row.get("pulled_s"))
        groups[run].append(Observation(
            sample_id=sample, run_id=run, x=seconds / 60 if seconds is not None else None,
            y=y, flag="missing" if y is None else "valid", source_id=source_id,
            excluded=excluded, exclusion_reason=reason,
            metadata={"mobile_sample": row, "source_row": index + 2},
        ))
        conditions[run] = {key: row.get(key) for key in (
            "process", "target", "oxidant", "oxidant_mm", "wavelength_nm",
            "fluence_rate_mw_cm2", "matrix", "matrix_lot", "matrix_ph",
        )}
        names[run] = _text(row.get("run_title")) or run
    return [Series(name=names[run], observations=points, y_unit="1", conditions=conditions[run],
                   source_ids=[source_id]) for run, points in groups.items()]


def _read_mobile_zip(data: bytes) -> tuple[dict, dict[str, bytes]]:
    files = {}
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000:
            raise ValueError("Mobile archive has too many entries")
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if (entry.filename in files or path.is_absolute() or ".." in path.parts
                    or "\\" in entry.filename or (entry.external_attr >> 16) & 0o170000 == 0o120000):
                raise ValueError("Unsafe or duplicate mobile archive path")
            if entry.is_dir():
                continue
            if (entry.filename != "workspace.json"
                    and not re.fullmatch(r"photos/[a-f0-9]{32}(\.thumb)?\.(jpg|png|webp)", entry.filename)):
                raise ValueError("Unexpected mobile archive member")
            total += entry.file_size
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("Mobile archive expands beyond 150 MiB")
            content = archive.read(entry)
            if len(content) != entry.file_size:
                raise ValueError("Invalid mobile archive member size")
            files[entry.filename] = content
    if "workspace.json" not in files:
        raise ValueError("Mobile archive is missing workspace.json")
    workspace = json.loads(files["workspace.json"])
    if not isinstance(workspace, dict):
        raise ValueError("Unsupported mobile workspace schema")
    lab = workspace.get("lab", workspace)
    if (not isinstance(lab, dict) or lab.get("version") not in (1, 2)
            or not isinstance(lab.get("experiments"), list)
            or not isinstance(lab.get("samples", []), list)):
        raise ValueError("Unsupported mobile workspace schema")
    experiment_ids = [e.get("id") if isinstance(e, dict) else None for e in lab["experiments"]]
    if (any(not isinstance(i, str) or not i.strip() for i in experiment_ids)
            or len(set(experiment_ids)) != len(experiment_ids)):
        raise ValueError("Invalid or duplicate mobile experiment identity")
    if any(not isinstance(sample, dict) for sample in lab.get("samples", [])):
        raise ValueError("Invalid mobile sample record")
    return workspace, files


def import_mobile(
    project: AnalysisProject, store: AnalysisStore, filename: str, data: bytes,
) -> list[Series]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("Upload exceeds 100 MiB")
    digest = hashlib.sha256(data).hexdigest()
    if any(source.sha256 == digest for source in project.sources):
        return []
    if Path(filename).suffix.lower() != ".zip":
        rows = parse_table(data, filename)
        series = _mobile_csv(rows, "")
        source = store.add_source(project, filename, data, "mobile")
    else:
        workspace, files = _read_mobile_zip(data)
        lab = workspace.get("lab", workspace)
        experiments = {e["id"]: e for e in lab["experiments"]}
        groups: dict[str, list[Observation]] = defaultdict(list)
        seen = set()
        for sample in lab.get("samples", []):
            run, sample_id = sample["experiment_id"], sample["id"]
            if run not in experiments or sample_id in seen:
                raise ValueError("Invalid or duplicate mobile sample identity")
            seen.add(sample_id)
            excluded = bool(sample.get("fit_excluded")) or bool(sample.get("archived"))
            reason = _text(sample.get("fit_exclusion_reason"))
            if sample.get("archived"):
                reason = reason or "archived_mobile_sample"
            if excluded and not reason:
                reason = "Imported exclusion; original reason not recorded"
            y = _number(sample.get("c_over_c0"))
            elapsed = _number(sample.get("elapsed_ms"))
            groups[run].append(Observation(
                sample_id=sample_id, run_id=run, x=elapsed / 60000 if elapsed is not None else None,
                y=y, flag="missing" if y is None else "valid", excluded=excluded,
                exclusion_reason=reason, metadata={"mobile_sample": sample},
            ))
        series = [Series(
            name=experiment.get("title", run), y_unit="1", observations=groups.get(run, []),
            conditions={"mobile_experiment": experiment, "reaction": experiment.get("run"),
                        "water": experiment.get("water"),
                        "time_origin": "mobile_experiment_start; analysis offset may be set"},
        ) for run, experiment in experiments.items()]
        source = store.add_source(project, filename, data, "mobile")
        base = f"mobile/{source.sha256}"
        for member, content in files.items():
            _atomic_write(store.project_dir(project.id) / base / member, content)
        project.mobile_archives.append({
            "source_id": source.id, "workspace": workspace,
            "members": {name: f"{base}/{name}" for name in files},
        })
    for item in series:
        item.source_ids = [source.id]
        for observation in item.observations:
            observation.source_id = source.id
    project.series.extend(series)
    project.history.append({"at": now(), "action": "import_mobile_snapshot",
                            "source_id": source.id, "series_ids": [s.id for s in series]})
    return series


def mobile_photos(project: AnalysisProject, series_id: str | None = None) -> list[dict[str, str]]:
    source_ids = None
    if series_id is not None:
        source_ids = next(s.source_ids for s in project.series if s.id == series_id)
    photos = []
    for archive in project.mobile_archives:
        if source_ids is not None and archive["source_id"] not in source_ids:
            continue
        for member in archive["members"]:
            if member.startswith("photos/") and ".thumb." not in member:
                photos.append({"source_id": archive["source_id"], "member": member,
                               "name": PurePosixPath(member).name})
    return photos


def attach_measurements(
    project: AnalysisProject, series_id: str, rows: list[dict[str, Any]], mapping: dict[str, str],
    *, y_unit: str, observable: str, reason: str, source_id: str = "",
) -> int:
    """Attach quantified assay results by sample ID, retaining phone sampling metadata."""
    if not reason.strip() or not mapping.get("sample_id") or not mapping.get("y"):
        raise ValueError("Linking requires sample ID, measurement, and a reason")
    series = next(s for s in project.series if s.id == series_id)
    lookup = {o.sample_id: o for o in series.observations}
    updates = []
    seen = set()
    for row in rows:
        sample = _text(row.get(mapping["sample_id"]))
        if sample not in lookup or sample in seen:
            raise ValueError(f"Unknown or duplicate sample ID: {sample}")
        seen.add(sample)
        raw = row.get(mapping["y"])
        flag = _flag(row.get(mapping.get("flag", "")), raw)
        value = None if flag in ("below_lod", "below_loq") and _text(raw).startswith("<") else _number(raw)
        sigma = _number(row.get(mapping.get("sigma", "")))
        before = lookup[sample]
        revised = Observation.model_validate({**before.model_dump(), "y": value, "sigma": sigma,
                                              "flag": flag})
        revised.metadata = {**before.metadata, "assay": {"source_id": source_id, "input": row}}
        updates.append((before, revised))
    unit_changed = _unit_key(y_unit) != _unit_key(series.y_unit)
    if unit_changed:
        if any(o.y is not None and o.sample_id not in seen for o in series.observations):
            raise ValueError("Changing assay units requires results for every previously measured sample")
    validated = Series.model_validate({**series.model_dump(), "y_unit": y_unit,
                                      "observable": observable})
    for before, revised in updates:
        index = next(i for i, o in enumerate(series.observations) if o.id == before.id)
        project.history.append({"at": now(), "action": "attach_measurement", "reason": reason,
                                "series_id": series_id, "before": before.model_dump(),
                                "after": revised.model_dump()})
        series.observations[index] = revised
    project.history.append({"at": now(), "action": "assay_units", "series_id": series_id,
                            "before": {"unit": series.y_unit, "observable": series.observable},
                            "after": {"unit": y_unit, "observable": observable}, "reason": reason})
    series.y_unit = validated.y_unit
    series.observable = validated.observable
    if unit_changed:
        previous_processing = series.processing.model_dump()
        series.processing = ProcessingConfig(
            target_x_unit=series.processing.target_x_unit,
            time_zero=series.processing.time_zero,
            aggregate_technical=series.processing.aggregate_technical,
        )
        project.history.append({
            "at": now(), "action": "reset_response_processing", "series_id": series_id,
            "reason": "Assay response units changed; review new response processing rules",
            "assay_reason": reason, "before": previous_processing,
            "after": series.processing.model_dump(),
        })
    if source_id and source_id not in series.source_ids:
        series.source_ids.append(source_id)
    invalidate_fits(project, series_id, reason)
    return len(updates)


def demo_project(store: AnalysisStore | None = None) -> AnalysisProject:
    """Self-created, explicitly synthetic demonstration; no personal records."""
    project = AnalysisProject(name="Synthetic UV/H2O2 experiment")
    project.series = [Series(
        name="Model compound · synthetic", conditions={
            "process": "UV/H2O2", "target": "Synthetic model compound",
            "pH": 7, "temperature_C": 25, "oxidant_mM": 1,
            "data_origin": "synthetic demonstration",
        }, observations=[Observation(
            x=t, y=10 * math.exp(-0.12 * t), run_id="synthetic-run-1",
            sample_id=f"synthetic-{i}", metadata={"synthetic": True},
        ) for i, t in enumerate([0, 2, 4, 6, 8, 12, 16])],
    )]
    if store:
        rows = [o.model_dump() for o in project.series[0].observations]
        source = store.add_source(project, "synthetic-input.json",
                                  json.dumps(rows, indent=2).encode("utf-8"), "manual")
        project.series[0].source_ids = [source.id]
        for observation in project.series[0].observations:
            observation.source_id = source.id
        store.save(project)
    return project
