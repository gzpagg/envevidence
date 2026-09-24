import csv
import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import Project


def safe_cell(value):
    """Spreadsheet text must not execute as a formula when opened in Excel."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx_cell(value):
    value = safe_cell(value)
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda m: f"[U+{ord(m[0]):04X}]", value)
        if len(value) > 32767:
            value = value[:32690] + " [TRUNCATED: full value in project JSON / CSV]"
    return value


def evidence_rows(project: Project) -> list[dict]:
    studies = {s.id: s for s in project.studies}
    rows = []
    for experiment in project.experiments:
        for field in experiment.fields:
            sources = field.current_sources
            rows.append(
                {
                    "study": studies[experiment.study_id].name,
                    "experiment": experiment.label,
                    "experiment_id": experiment.id,
                    "field_id": field.id,
                    "field": field.key,
                    "label": field.label,
                    "model_value": field.original.value,
                    "model_unit": field.original.unit,
                    "value": field.value,
                    "unit": field.unit,
                    "extraction_status": field.original.status,
                    "human_review": field.review_status,
                    "source_located": bool(sources) and all(s.matched for s in sources),
                    "sources": json.dumps([s.model_dump() for s in sources], ensure_ascii=False),
                    "original_issues": " | ".join(field.issues),
                    "revision_count": len(field.revisions),
                    "provider": project.provider,
                    "model": project.model,
                    "prompt_version": project.prompt_version,
                }
            )
    return rows


def csv_bytes(project: Project) -> bytes:
    rows = evidence_rows(project)
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(
        stream, fieldnames=list(rows[0]) if rows else ["study", "experiment", "field", "value"]
    )
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe_cell(v) for k, v in row.items()})
    return stream.getvalue().encode("utf-8-sig")


def xlsx_bytes(project: Project) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, rows, empty_headers):
        ws = workbook.create_sheet(name)
        headers = list(rows[0]) if rows else empty_headers
        ws.append(headers)
        for row in rows:
            ws.append([xlsx_cell(row.get(k)) for k in headers])
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="147D73")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(
                55, max(18, len(str(column[0].value)) + 3)
            )

    sheet("Evidence", evidence_rows(project), ["study", "field", "value"])
    study_names = {s.id: s.name for s in project.studies}
    wide = []
    for exp in project.experiments:
        row = {"study": study_names[exp.study_id], "experiment": exp.label}
        for field in exp.fields:
            row[field.key] = field.value
            row[field.key + "_unit"] = field.unit
            row[field.key + "_review"] = field.review_status
        wide.append(row)
    sheet("Experiments", wide, ["study", "experiment"])
    revisions = []
    for exp in project.experiments:
        for field in exp.fields:
            for revision in field.revisions:
                row = {
                    "field_id": field.id,
                    "experiment": exp.label,
                    "field": field.key,
                    **revision.model_dump(exclude={"sources"}),
                }
                row["sources"] = json.dumps(
                    [s.model_dump() for s in revision.sources], ensure_ascii=False
                )
                revisions.append(row)
    sheet("Revisions", revisions, ["field_id", "timestamp", "reviewer", "note"])
    documents = [
        {
            "study": s.name,
            "file": d.filename,
            "role": d.role,
            "sha256": d.sha256,
            "pages": d.page_count,
            "parsed_pages": len(d.blocks),
            "parser": d.parser,
            "warnings": " | ".join(d.warnings),
        }
        for s in project.studies
        for d in s.documents
    ]
    sheet("Documents", documents, ["study", "file", "sha256"])
    sheet(
        "Runs",
        [r.model_dump(exclude={"usage"}) | {"usage": json.dumps(r.usage)} for r in project.runs],
        ["study_id", "status"],
    )
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
